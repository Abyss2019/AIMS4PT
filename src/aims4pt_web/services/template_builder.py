"""Build the downloadable AIMS4PT input template in memory."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from aims4pt_web.services.input_validation import (
    CPX_OXIDE_COLUMNS,
    EXPECTED_INPUT_COLUMNS,
    LIQUID_OXIDE_COLUMNS,
)


def build_input_template() -> BytesIO:
    """Return an XLSX template with grouped headers and example rows."""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "compositions"

    group_row = (
        ["Sample metadata"]
        + ["Clinopyroxene oxide wt%"] * len(CPX_OXIDE_COLUMNS)
        + ["Liquid oxide wt%"] * len(LIQUID_OXIDE_COLUMNS)
    )
    worksheet.append(group_row)
    worksheet.append(EXPECTED_INPUT_COLUMNS)
    worksheet.append(
        [
            "sample-1",
            51.2,
            0.7,
            5.1,
            6.8,
            0.2,
            15.1,
            21.0,
            0.35,
            0.01,
            0.25,
            50.1,
            1.1,
            16.8,
            8.2,
            0.15,
            6.5,
            9.8,
            3.2,
            1.4,
            0.03,
            0.2,
            1.0,
        ]
    )
    worksheet.append(
        [
            "sample-2",
            50.5,
            0.6,
            4.8,
            7.1,
            0.18,
            14.8,
            21.4,
            0.32,
            0.02,
            0.18,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
        ]
    )

    header_fill = PatternFill("solid", fgColor="E9ECEF")
    group_fill = PatternFill("solid", fgColor="DDE7F0")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = group_fill
        cell.alignment = Alignment(horizontal="center")
    for cell in worksheet[2]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=1)
    worksheet.merge_cells(
        start_row=1,
        start_column=2,
        end_row=1,
        end_column=1 + len(CPX_OXIDE_COLUMNS),
    )
    worksheet.merge_cells(
        start_row=1,
        start_column=2 + len(CPX_OXIDE_COLUMNS),
        end_row=1,
        end_column=len(EXPECTED_INPUT_COLUMNS),
    )
    worksheet.freeze_panes = "A3"

    for idx, column in enumerate(EXPECTED_INPUT_COLUMNS, start=1):
        width = max(12, len(column) + 2)
        worksheet.column_dimensions[get_column_letter(idx)].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output

