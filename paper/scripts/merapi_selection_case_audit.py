"""Audit Merapi sample-level AIMS4PT_cpx model-selection patterns.

This script reads existing Merapi workflow reports and cached cpx-liquid
pairing files. It does not retrain resources, rebuild thermobarometer
workflows, or rerun sample pairing. The audit classifies model-selection
patterns, not geological reservoirs or compositional clusters.
"""

from __future__ import annotations

import argparse
import itertools
import math
import re
import sys
import textwrap
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

try:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
except Exception:  # pragma: no cover - Pillow chart path is used below.
    matplotlib = None
    plt = None
import numpy as np
import pandas as pd

try:
    from constants_illustration import get_model_abbreviation
except Exception:  # pragma: no cover - fallback for package-style execution.
    try:
        from paper.scripts.constants_illustration import get_model_abbreviation
    except Exception:  # pragma: no cover
        get_model_abbreviation = None


# -----------------------------
# Configurable audit thresholds
# -----------------------------

MIN_PROP_RETAIN = 0.05
DOMINANT_P1 = 0.683
NEAR_DOMINANT_P1 = 0.60
GAP_DOMINANT = 0.25
TWO_MODEL_COVERAGE = 0.80
HIGH_DISPERSION_NEFF = 2.5
HIGHLY_DISPERSED_P1 = 0.40

P_LOW_IMPACT_KBAR = 0.5
P_HIGH_IMPACT_KBAR = 1.0
T_LOW_IMPACT_C = 20.0
T_HIGH_IMPACT_C = 50.0

USE_TUKEY_FOR_PT = True
DEFAULT_TIMESTAMP = "add_pre-2006_028"
OUTPUT_DIR = Path("paper/outputs/merapi_selection_audit")

MOLAR_MASS_MGO = 40.3044
MOLAR_MASS_FEO = 71.844

BRANCHES = [
    ("P", "cpx_only", "2006"),
    ("P", "cpx_only", "2010"),
    ("P", "cpx_liq", "2006"),
    ("P", "cpx_liq", "2010"),
    ("T", "cpx_only", "2006"),
    ("T", "cpx_only", "2010"),
    ("T", "cpx_liq", "2006"),
    ("T", "cpx_liq", "2010"),
]

EXPECTED_BEHAVIOR = {
    ("P", "cpx_only", "2006"): "Expected dominant or near-dominant; Jor22 should be the main model.",
    ("T", "cpx_only", "2006"): "Expected ambiguous unless composition separation is unexpectedly strong; Hig21 and Pet20 should compete.",
    ("T", "cpx_liq", "2006"): "Expected near-dominant or low-impact mixed; Chi23 may dominate but AgL24 may give similar temperatures.",
    ("P", "cpx_liq", "2006"): "Expected Pet20-leading but not necessarily dominant; Pet20-Jor22 pressure impact and composition structure need checking.",
    ("P", "cpx_only", "2010"): "Expected near-dominant or two-model competition; Pet20 likely leads with Jor22 secondary.",
    ("T", "cpx_only", "2010"): "Expected composition-associated or mixed if AgL24-selected samples occupy a high-Al2O3/low-Mg# cpx domain.",
    ("T", "cpx_liq", "2010"): "Expected low- to moderate-impact mixed if AgL24, Chi23, and Jor22 have similar median temperatures.",
    ("P", "cpx_liq", "2010"): "Expected dispersed/model-dependent; high-pressure tail may depend on low-proportion Chi23 or AgL24 support.",
}

COLUMN_ALIASES = {
    "sample_id": ["Sample_id", "Sample ID", "sample_id", "sample"],
    "crystal_id": ["Crystal_id", "Crystal ID", "crystal_id"],
    "eruption": ["Eruption", "eruption", "eruption_year", "year"],
    "selected_model": ["Selected_model", "selected_model", "favored_model", "best_model"],
    "cpx_sio2": ["cpx_SiO2", "SiO2", "mine__SiO2"],
    "cpx_al2o3": ["cpx_Al2O3", "Al2O3", "mine__Al2O3"],
    "cpx_na2o": ["cpx_Na2O", "Na2O", "mine__Na2O"],
    "cpx_cao": ["cpx_CaO", "CaO", "mine__CaO"],
    "cpx_mgo": ["cpx_MgO", "MgO", "mine__MgO"],
    "cpx_feo": ["cpx_FeO", "FeO", "mine__FeO"],
    "cpx_feot": ["cpx_FeOt", "FeOt", "mine__FeOt"],
    "liq_sio2": ["liq_SiO2", "liq__SiO2", "SiO2_liq"],
    "liq_na2o": ["liq_Na2O", "liq__Na2O", "Na2O_liq"],
    "liq_k2o": ["liq_K2O", "liq__K2O", "K2O_liq"],
    "liq_mgo": ["liq_MgO", "liq__MgO", "MgO_liq"],
    "liq_feo": ["liq_FeO", "liq__FeO", "FeO_liq"],
    "liq_feot": ["liq_FeOt", "liq__FeOt", "FeOt_liq"],
    "liq_h2o": ["liq_H2O", "liq__H2O", "H2O_liq"],
}


@dataclass(frozen=True)
class BranchKey:
    target: str
    phase_type: str
    year: str

    @property
    def phase_label(self) -> str:
        return self.phase_type.replace("_", "-")

    @property
    def label(self) -> str:
        return f"{self.target} {self.phase_label} {self.year}"

    @property
    def stem_token(self) -> str:
        return f"marapi_{self.year}_{self.target}_workflow_{self.phase_type}_{self.target}"


@dataclass
class BranchData:
    key: BranchKey
    report_path: Path
    frame: pd.DataFrame
    model_names: list[str]
    retained_models: list[str]
    support: dict[str, Any]
    summary_file_rows: pd.DataFrame


def configure_stdout() -> None:
    """Use UTF-8 console output when Python supports reconfiguration."""
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass


def warn(warnings_out: list[dict[str, str]], message: str, *, branch: str = "", file: Path | None = None) -> None:
    """Print and store a diagnostic warning."""
    prefix = "[WARN]"
    if branch:
        prefix += f" [{branch}]"
    if file is not None:
        prefix += f" [{file}]"
    print(f"{prefix} {message}")
    warnings_out.append({"branch": branch, "file": str(file) if file else "", "warning": message})


def repo_root_from_script() -> Path:
    """Return the repository root when the script is run from any location."""
    return Path(__file__).resolve().parents[2]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repo-root", type=Path, default=repo_root_from_script())
    parser.add_argument("--timestamp", default=DEFAULT_TIMESTAMP)
    parser.add_argument("--cache-dir", type=Path, default=None)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    parser.add_argument(
        "--no-tukey",
        action="store_true",
        help="Use all finite P-T values instead of recomputing the 1.5 IQR Tukey fence.",
    )
    return parser.parse_args()


def find_cache_dir(repo_root: Path, timestamp: str, explicit_cache_dir: Path | None, warnings_out: list[dict[str, str]]) -> Path:
    """Locate the Merapi cache directory containing the eight report files."""
    if explicit_cache_dir is not None:
        cache_dir = explicit_cache_dir if explicit_cache_dir.is_absolute() else repo_root / explicit_cache_dir
        if not cache_dir.exists():
            raise FileNotFoundError(f"Explicit cache directory does not exist: {cache_dir}")
        return cache_dir

    preferred = repo_root / "paper" / ".cache" / timestamp
    if preferred.exists():
        return preferred

    cache_root = repo_root / "paper" / ".cache"
    candidates = []
    for child in cache_root.glob("add_pre-2006*"):
        if not child.is_dir():
            continue
        count = sum(1 for key in (BranchKey(*b) for b in BRANCHES) if list(child.glob(f"{key.stem_token}*_report.xlsx")))
        if count:
            candidates.append((count, child.stat().st_mtime, child))

    if not candidates:
        raise FileNotFoundError(f"No Merapi cache directory with report files was found under {cache_root}")

    candidates.sort(key=lambda item: (item[0], item[1]), reverse=True)
    chosen = candidates[0][2]
    warn(warnings_out, f"Preferred timestamp {timestamp!r} was not found; using {chosen}.")
    return chosen


def find_report_file(cache_dir: Path, key: BranchKey) -> Path:
    """Find one report file for a branch."""
    pattern = f"{key.stem_token}*_report.xlsx"
    matches = sorted(cache_dir.glob(pattern))
    if not matches:
        raise FileNotFoundError(f"Missing report for {key.label}: {cache_dir / pattern}")
    if len(matches) > 1:
        matches.sort(key=lambda path: path.stat().st_mtime, reverse=True)
    return matches[0]


def find_summary_file(cache_dir: Path) -> Path | None:
    path = cache_dir / "merapi_fig8_9_model_summary.xlsx"
    return path if path.exists() else None


def find_fig9_data_file(cache_dir: Path) -> Path | None:
    path = cache_dir / "fig_9_Merapi_different_constraints_v2_data.xlsx"
    return path if path.exists() else None


def find_pairing_file(cache_dir: Path, year: str) -> Path | None:
    matches = sorted(cache_dir.glob(f"merapi_{year}_pairing*k*d028*endmembers.xlsx"))
    if not matches:
        matches = sorted(cache_dir.glob(f"merapi_{year}_pairing*endmembers.xlsx"))
    return matches[0] if matches else None


def clean_text(value: Any) -> str:
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return ""
    return str(value).strip()


def flatten_results_columns(columns: pd.MultiIndex) -> list[str]:
    """Flatten report Results columns while preserving model names."""
    flat: list[str] = []
    for group, name in columns:
        group_s = clean_text(group)
        name_s = clean_text(name)
        if group_s.startswith("Unnamed") and name_s.startswith("Unnamed"):
            flat.append("_row_id")
        elif group_s == "Original Data":
            flat.append(name_s)
        elif group_s == "Model Predictions":
            flat.append(f"prediction::{name_s}")
        else:
            flat.append(name_s or group_s)
    return flat


def read_results_sheet(report_path: Path, warnings_out: list[dict[str, str]], key: BranchKey) -> tuple[pd.DataFrame, list[str]]:
    """Read the Results sheet from an AIMS4PT report workbook."""
    df = pd.read_excel(report_path, sheet_name="Results", header=[0, 1])
    df.columns = flatten_results_columns(df.columns)
    df = df.dropna(how="all").copy()
    model_names = [col.split("prediction::", 1)[1] for col in df.columns if col.startswith("prediction::")]
    if "Sample_id" in df.columns:
        df = df[df["Sample_id"].notna()].copy()
    elif model_names:
        prediction_cols = [f"prediction::{name}" for name in model_names]
        df = df[df[prediction_cols].notna().any(axis=1)].copy()
        warn(warnings_out, "Results sheet lacks Sample_id; retained rows with any model prediction.", branch=key.label, file=report_path)
    else:
        warn(warnings_out, "Results sheet lacks model prediction columns.", branch=key.label, file=report_path)
    df = df.reset_index(drop=True)
    return df, model_names


def read_ranking_sheet(report_path: Path, warnings_out: list[dict[str, str]], key: BranchKey) -> pd.DataFrame:
    """Read per-sample deviation rankings and selected models."""
    rank = pd.read_excel(report_path, sheet_name="Ranking Details")
    rank = rank.dropna(how="all").copy()
    first_col = rank.columns[0]
    if str(first_col).startswith("Unnamed"):
        rank = rank.rename(columns={first_col: "_row_id"})
    if "Selected_model" not in rank.columns:
        alt = first_existing_column(rank, COLUMN_ALIASES["selected_model"])
        if alt is not None:
            rank = rank.rename(columns={alt: "Selected_model"})
        else:
            warn(warnings_out, "Ranking Details lacks Selected_model.", branch=key.label, file=report_path)
            rank["Selected_model"] = np.nan
    rank = rank.reset_index(drop=True)
    return rank


def first_existing_column(df: pd.DataFrame, candidates: Iterable[str]) -> str | None:
    """Return the first existing column from a list of aliases."""
    column_map = {str(col).casefold(): str(col) for col in df.columns}
    for candidate in candidates:
        if candidate in df.columns:
            return candidate
        found = column_map.get(str(candidate).casefold())
        if found is not None:
            return found
    return None


def numeric_series(df: pd.DataFrame, col: str | None) -> pd.Series:
    """Return a numeric series, or all-NaN when the column is missing."""
    if col is None or col not in df.columns:
        return pd.Series(np.nan, index=df.index, dtype=float)
    return pd.to_numeric(df[col], errors="coerce")


def calculate_mg_number(df: pd.DataFrame, mg_col: str | None, feo_col: str | None, feot_col: str | None) -> tuple[pd.Series, str]:
    """Calculate molar Mg# from MgO and FeO/FeOt."""
    mg = numeric_series(df, mg_col) / MOLAR_MASS_MGO
    fe_source = feo_col if feo_col is not None and feo_col in df.columns else None
    fe = numeric_series(df, fe_source) / MOLAR_MASS_FEO
    if fe.isna().all() and feot_col is not None and feot_col in df.columns:
        fe_source = feot_col
        fe = numeric_series(df, fe_source) / MOLAR_MASS_FEO
    denom = mg + fe
    mg_number = pd.Series(np.where(denom > 0, mg / denom, np.nan), index=df.index, dtype=float)
    return mg_number, clean_text(fe_source) if fe_source else ""


def add_canonical_composition_columns(df: pd.DataFrame, key: BranchKey, warnings_out: list[dict[str, str]]) -> pd.DataFrame:
    """Add canonical cpx/liquid composition columns used by the audit."""
    out = df.copy()
    cpx_map = {
        "cpx_SiO2": first_existing_column(out, COLUMN_ALIASES["cpx_sio2"]),
        "cpx_Al2O3": first_existing_column(out, COLUMN_ALIASES["cpx_al2o3"]),
        "cpx_Na2O": first_existing_column(out, COLUMN_ALIASES["cpx_na2o"]),
        "cpx_CaO": first_existing_column(out, COLUMN_ALIASES["cpx_cao"]),
        "cpx_MgO": first_existing_column(out, COLUMN_ALIASES["cpx_mgo"]),
    }
    for canonical, source in cpx_map.items():
        out[canonical] = numeric_series(out, source)
        if source is None:
            warn(warnings_out, f"Missing composition column for {canonical}.", branch=key.label)

    cpx_mg, cpx_fe_source = calculate_mg_number(
        out,
        first_existing_column(out, COLUMN_ALIASES["cpx_mgo"]),
        first_existing_column(out, COLUMN_ALIASES["cpx_feo"]),
        first_existing_column(out, COLUMN_ALIASES["cpx_feot"]),
    )
    out["cpx_MgNumber"] = cpx_mg
    out["cpx_MgNumber_Fe_source"] = cpx_fe_source

    if key.phase_type == "cpx_liq":
        liq_map = {
            "liq_SiO2": first_existing_column(out, COLUMN_ALIASES["liq_sio2"]),
            "liq_Na2O": first_existing_column(out, COLUMN_ALIASES["liq_na2o"]),
            "liq_K2O": first_existing_column(out, COLUMN_ALIASES["liq_k2o"]),
            "liq_MgO": first_existing_column(out, COLUMN_ALIASES["liq_mgo"]),
            "liq_H2O": first_existing_column(out, COLUMN_ALIASES["liq_h2o"]),
        }
        for canonical, source in liq_map.items():
            out[canonical] = numeric_series(out, source)
            if source is None and canonical != "liq_H2O":
                warn(warnings_out, f"Missing liquid composition column for {canonical}.", branch=key.label)
        out["liq_TotalAlkali"] = out["liq_Na2O"] + out["liq_K2O"]
        liq_mg, liq_fe_source = calculate_mg_number(
            out,
            first_existing_column(out, COLUMN_ALIASES["liq_mgo"]),
            first_existing_column(out, COLUMN_ALIASES["liq_feo"]),
            first_existing_column(out, COLUMN_ALIASES["liq_feot"]),
        )
        out["liq_MgNumber"] = liq_mg
        out["liq_MgNumber_Fe_source"] = liq_fe_source

    return out


def append_pairing_liquids(
    df: pd.DataFrame,
    cache_dir: Path,
    key: BranchKey,
    warnings_out: list[dict[str, str]],
) -> pd.DataFrame:
    """Append cached liquid compositions to cpx-liquid report rows."""
    if key.phase_type != "cpx_liq":
        return df

    pair_path = find_pairing_file(cache_dir, key.year)
    if pair_path is None:
        warn(warnings_out, "No cpx-liquid pairing workbook found; liquid composition metrics will be incomplete.", branch=key.label)
        return df

    pairs = pd.read_excel(pair_path, sheet_name="pairs")
    if "if_pass_test" in pairs.columns:
        pass_mask = pairs["if_pass_test"].astype(str).str.casefold().isin(["true", "1", "yes"])
        pairs = pairs.loc[pass_mask].copy()
    pairs = pairs.reset_index(drop=True)

    out = df.copy().reset_index(drop=True)
    sample_col = first_existing_column(out, COLUMN_ALIASES["sample_id"])
    pair_sample_col = "mine__Sample_id" if "mine__Sample_id" in pairs.columns else first_existing_column(pairs, ["Sample_id"])
    aligned = len(out) == len(pairs)
    if aligned and sample_col is not None and pair_sample_col is not None:
        left = out[sample_col].astype(str).str.strip().reset_index(drop=True)
        right = pairs[pair_sample_col].astype(str).str.strip().reset_index(drop=True)
        aligned = bool(left.equals(right))

    pair_cols = [
        col
        for col in pairs.columns
        if str(col).startswith("liq__") or str(col) in {"kd_value", "kd_error", "if_pass_test", "mine__pair_row_idx", "liq__pair_row_idx"}
    ]
    if aligned:
        for col in pair_cols:
            out[col] = pairs[col].to_numpy()
        out["pairing_file"] = str(pair_path)
        out["pairing_alignment"] = "row_order"
        return out

    warn(
        warnings_out,
        "Pairing rows did not align by count/sample order; trying Sample_id merge.",
        branch=key.label,
        file=pair_path,
    )
    if sample_col is None or pair_sample_col is None:
        warn(warnings_out, "Cannot merge pairing rows because Sample_id is missing.", branch=key.label, file=pair_path)
        return out

    merge_cols = [pair_sample_col] + pair_cols
    pair_small = pairs[merge_cols].copy()
    pair_small["_sample_key"] = pair_small[pair_sample_col].astype(str).str.strip()
    out["_sample_key"] = out[sample_col].astype(str).str.strip()
    out = out.merge(pair_small.drop(columns=[pair_sample_col]), on="_sample_key", how="left", validate="m:1")
    out = out.drop(columns=["_sample_key"])
    out["pairing_file"] = str(pair_path)
    out["pairing_alignment"] = "sample_id_merge"
    return out


def merge_results_and_ranking(results: pd.DataFrame, ranking: pd.DataFrame, key: BranchKey, warnings_out: list[dict[str, str]]) -> pd.DataFrame:
    """Merge Results and Ranking Details by row id when possible."""
    results = results.copy()
    ranking = ranking.copy()
    if "_row_id" in results.columns and "_row_id" in ranking.columns:
        results["_row_id"] = pd.to_numeric(results["_row_id"], errors="coerce")
        ranking["_row_id"] = pd.to_numeric(ranking["_row_id"], errors="coerce")
        merged = results.merge(ranking, on="_row_id", how="left", suffixes=("", "__rank"))
    else:
        if len(results) != len(ranking):
            warn(warnings_out, f"Results rows ({len(results)}) and Ranking rows ({len(ranking)}) differ; using row-order join.", branch=key.label)
        merged = pd.concat([results.reset_index(drop=True), ranking.reset_index(drop=True)], axis=1)

    if "Selected_model" not in merged.columns:
        warn(warnings_out, "Merged branch table lacks Selected_model.", branch=key.label)
        merged["Selected_model"] = np.nan
    merged["Selected_model"] = merged["Selected_model"].astype(object)
    return merged.reset_index(drop=True)


def model_abbrev(model_name: str, target: str) -> str:
    """Return a compact model label for tables and plots."""
    if model_name == "Other":
        return "Other"
    if get_model_abbreviation is not None:
        try:
            return str(get_model_abbreviation(str(model_name), target))
        except Exception:
            pass
    normalized = str(model_name)
    fallback = {
        "Petrelli": "Pet20",
        "Higgins": "Hig21",
        "Jorgenson": "Jor22",
        "Chicchi": "Chi23",
        "Wang": "Wan21",
        "Neave": "NP17",
        "Putirka": "Pu08",
    }
    if "greda" in normalized.casefold() or "agreda" in normalized.casefold():
        return "AgL24"
    for token, short in fallback.items():
        if token.casefold() in normalized.casefold():
            return short
    return normalized[:18]


def read_summary_rows(summary_path: Path | None, key: BranchKey) -> pd.DataFrame:
    """Read existing Fig. 8-9 summary rows for one branch when available."""
    if summary_path is None or not summary_path.exists():
        return pd.DataFrame()
    sheet = f"{key.target}_summary"
    try:
        summary = pd.read_excel(summary_path, sheet_name=sheet)
    except Exception:
        return pd.DataFrame()
    mask = (
        summary["eruption"].astype(str).eq(key.year)
        & summary["phase_type"].astype(str).eq(key.phase_type)
        & summary["quantity"].astype(str).eq(key.target)
    )
    return summary.loc[mask].copy()


def read_branch_data(cache_dir: Path, summary_path: Path | None, key: BranchKey, warnings_out: list[dict[str, str]]) -> BranchData:
    """Read and assemble all per-sample data needed for one branch."""
    report_path = find_report_file(cache_dir, key)
    results, model_names = read_results_sheet(report_path, warnings_out, key)
    ranking = read_ranking_sheet(report_path, warnings_out, key)
    frame = merge_results_and_ranking(results, ranking, key, warnings_out)
    frame = append_pairing_liquids(frame, cache_dir, key, warnings_out)
    frame = add_canonical_composition_columns(frame, key, warnings_out)
    frame["target"] = key.target
    frame["phase_type"] = key.phase_type
    frame["model_type"] = key.phase_label
    frame["eruption_year"] = key.year
    frame["branch"] = key.label
    frame["report_file"] = str(report_path)
    summary_file_rows = read_summary_rows(summary_path, key)
    support = compute_support_metrics(frame, key)
    retained_models = support["retained_models"]
    return BranchData(key, report_path, frame, model_names, retained_models, support, summary_file_rows)


def detect_input_files(cache_dir: Path) -> list[Path]:
    """Return relevant input files for diagnostics."""
    files: list[Path] = []
    files.extend(sorted(cache_dir.glob("*report.xlsx")))
    for name in [
        "merapi_fig8_9_model_summary.xlsx",
        "fig_9_Merapi_different_constraints_v2_data.xlsx",
        "merapi_2006_pairing_kd028_err1e-4_n20000_endmembers.xlsx",
        "merapi_2010_pairing_kd028_err1e-4_n20000_endmembers.xlsx",
    ]:
        path = cache_dir / name
        if path.exists():
            files.append(path)
    return list(dict.fromkeys(files))


def diagnostic_sheet_columns(path: Path, sheet_name: str) -> list[str]:
    """Read a compact column listing for diagnostics."""
    if sheet_name == "Results":
        df = pd.read_excel(path, sheet_name=sheet_name, header=[0, 1], nrows=2)
        return flatten_results_columns(df.columns)
    df = pd.read_excel(path, sheet_name=sheet_name, nrows=2)
    return [str(col) for col in df.columns]


def print_input_diagnostics(files: list[Path]) -> pd.DataFrame:
    """Print detected input files, sheet names, and column names."""
    rows: list[dict[str, Any]] = []
    print("\nDetected Merapi input files and sheet columns")
    print("--------------------------------------------")
    for path in files:
        print(f"\nFILE: {path}")
        try:
            xls = pd.ExcelFile(path)
        except Exception as exc:
            print(f"  ERROR opening workbook: {exc}")
            rows.append({"file": str(path), "sheet": "", "columns": "", "error": str(exc)})
            continue
        print(f"  sheets: {xls.sheet_names}")
        for sheet in xls.sheet_names:
            try:
                cols = diagnostic_sheet_columns(path, sheet)
                print(f"  SHEET: {sheet}")
                print(f"    columns: {cols}")
                rows.append({"file": str(path), "sheet": sheet, "columns": " | ".join(cols), "error": ""})
            except Exception as exc:
                print(f"  SHEET: {sheet} ERROR: {exc}")
                rows.append({"file": str(path), "sheet": sheet, "columns": "", "error": str(exc)})
    return pd.DataFrame(rows)


def compute_support_metrics(df: pd.DataFrame, key: BranchKey) -> dict[str, Any]:
    """Compute favored-model support concentration metrics."""
    selected = df["Selected_model"].dropna().astype(str)
    selected = selected[selected.str.len() > 0]
    n = int(selected.size)
    if n == 0:
        return {
            "N": 0,
            "retained_models": [],
            "props": pd.Series(dtype=float),
            "grouped_props": pd.Series(dtype=float),
        }

    props = selected.value_counts(normalize=True)
    counts = selected.value_counts()
    retained = props[props >= MIN_PROP_RETAIN].copy()
    other_prop = float(props[props < MIN_PROP_RETAIN].sum())
    grouped = retained.copy()
    if other_prop > 0:
        grouped.loc["Other"] = other_prop
    grouped = grouped.sort_values(ascending=False)

    sorted_props = props.sort_values(ascending=False)
    retained_sorted = retained.sort_values(ascending=False)
    p = [float(retained_sorted.iloc[i]) if i < len(retained_sorted) else 0.0 for i in range(4)]
    cumulative = sorted_props.cumsum()

    def k_for_threshold(threshold: float) -> float:
        hits = np.flatnonzero(cumulative.to_numpy() >= threshold)
        return float(hits[0] + 1) if hits.size else np.nan

    retained_norm = retained_sorted / retained_sorted.sum() if retained_sorted.sum() > 0 else retained_sorted
    k = int(len(retained_sorted))
    if k == 0:
        n_eff = np.nan
        h_norm = np.nan
    else:
        n_eff = float(1.0 / np.sum(np.square(retained_norm.to_numpy(dtype=float))))
        h_norm = 0.0 if k == 1 else float(-np.sum(retained_norm * np.log(retained_norm)) / math.log(k))

    return {
        "N": n,
        "K": k,
        "props": props,
        "counts": counts,
        "retained_models": retained_sorted.index.tolist(),
        "retained_support_sum": float(retained_sorted.sum()),
        "grouped_props": grouped,
        "other_prop": other_prop,
        "p1": p[0],
        "p2": p[1],
        "p3": p[2],
        "p4": p[3],
        "gap12": p[0] - p[1],
        "C2": float(np.sum(p[:2])),
        "C3": float(np.sum(p[:3])),
        "C4": float(np.sum(p[:4])),
        "K80": k_for_threshold(0.80),
        "K90": k_for_threshold(0.90),
        "N_eff": n_eff,
        "H_norm": h_norm,
    }


def model_distribution_rows(branch: BranchData) -> list[dict[str, Any]]:
    """Return model support rows for actual and grouped proportions."""
    rows: list[dict[str, Any]] = []
    props: pd.Series = branch.support.get("props", pd.Series(dtype=float))
    counts: pd.Series = branch.support.get("counts", pd.Series(dtype=int))
    grouped: pd.Series = branch.support.get("grouped_props", pd.Series(dtype=float))
    for rank, (model, prop) in enumerate(props.sort_values(ascending=False).items(), start=1):
        rows.append(
            {
                "branch": branch.key.label,
                "target": branch.key.target,
                "phase_type": branch.key.phase_type,
                "model_type": branch.key.phase_label,
                "eruption_year": branch.key.year,
                "row_type": "actual_model",
                "rank": rank,
                "model": model,
                "model_abbrev": model_abbrev(model, branch.key.target),
                "n_samples": int(counts.get(model, 0)),
                "favored_proportion": float(prop),
                "retained": bool(prop >= MIN_PROP_RETAIN),
            }
        )
    for rank, (model, prop) in enumerate(grouped.items(), start=1):
        rows.append(
            {
                "branch": branch.key.label,
                "target": branch.key.target,
                "phase_type": branch.key.phase_type,
                "model_type": branch.key.phase_label,
                "eruption_year": branch.key.year,
                "row_type": "grouped_for_support",
                "rank": rank,
                "model": model,
                "model_abbrev": model_abbrev(model, branch.key.target),
                "n_samples": int(round(float(prop) * branch.support.get("N", 0))),
                "favored_proportion": float(prop),
                "retained": model != "Other",
            }
        )
    return rows


def finite_values(values: Any) -> np.ndarray:
    """Return finite numeric values from an array-like object."""
    arr = pd.to_numeric(pd.Series(values), errors="coerce").replace([np.inf, -np.inf], np.nan)
    arr = arr.dropna().to_numpy(dtype=float)
    return arr[np.isfinite(arr)]


def tukey_filter(values: np.ndarray) -> np.ndarray:
    """Apply a 1.5 IQR Tukey fence, matching the Merapi helper."""
    arr = np.asarray(values, dtype=float).ravel()
    arr = arr[np.isfinite(arr)]
    if arr.size <= 3:
        return arr
    q1, q3 = np.nanpercentile(arr, [25, 75])
    iqr = q3 - q1
    if not np.isfinite(iqr) or iqr == 0:
        return arr
    lower = q1 - 1.5 * iqr
    upper = q3 + 1.5 * iqr
    return arr[(arr >= lower) & (arr <= upper)]


def summarize_numeric_values(values: Any, *, use_tukey: bool) -> dict[str, Any]:
    """Return robust distribution summary statistics."""
    raw = finite_values(values)
    used = tukey_filter(raw) if use_tukey else raw
    out = {
        "n_raw": int(raw.size),
        "n_valid": int(used.size),
        "filter": "tukey_fence_recomputed" if use_tukey else "all_finite_values",
        "mean": np.nan,
        "median": np.nan,
        "q25": np.nan,
        "q75": np.nan,
        "iqr": np.nan,
        "min": np.nan,
        "max": np.nan,
    }
    if used.size == 0:
        return out
    q25, median, q75 = np.nanpercentile(used, [25, 50, 75])
    out.update(
        {
            "mean": float(np.nanmean(used)),
            "median": float(median),
            "q25": float(q25),
            "q75": float(q75),
            "iqr": float(q75 - q25),
            "min": float(np.nanmin(used)),
            "max": float(np.nanmax(used)),
        }
    )
    return out


def prediction_column(df: pd.DataFrame, model_name: str) -> str | None:
    col = f"prediction::{model_name}"
    if col in df.columns:
        return col
    for candidate in df.columns:
        if str(candidate).startswith("prediction::") and str(candidate).split("prediction::", 1)[1] == str(model_name):
            return str(candidate)
    return None


def compute_pt_distribution_rows(branch: BranchData, *, use_tukey: bool, warnings_out: list[dict[str, str]]) -> list[dict[str, Any]]:
    """Summarize retained model P-T distributions for all and favored subsets."""
    rows: list[dict[str, Any]] = []
    df = branch.frame
    for model in branch.retained_models:
        pred_col = prediction_column(df, model)
        if pred_col is None:
            warn(warnings_out, f"Missing prediction column for retained model {model}.", branch=branch.key.label)
            continue
        scopes = {
            "all_predictions": df[pred_col],
            "favored_subset": df.loc[df["Selected_model"].astype(str).eq(model), pred_col],
        }
        for scope, values in scopes.items():
            stats = summarize_numeric_values(values, use_tukey=use_tukey)
            rows.append(
                {
                    "branch": branch.key.label,
                    "target": branch.key.target,
                    "unit": "kbar" if branch.key.target == "P" else "degC",
                    "phase_type": branch.key.phase_type,
                    "model_type": branch.key.phase_label,
                    "eruption_year": branch.key.year,
                    "model": model,
                    "model_abbrev": model_abbrev(model, branch.key.target),
                    "scope": scope,
                    **stats,
                }
            )
    return rows


def impact_level(target: str, delta_median: float) -> str:
    """Classify P-T distribution impact from median spread."""
    if not np.isfinite(delta_median):
        return "unknown"
    if target == "P":
        if delta_median <= P_LOW_IMPACT_KBAR:
            return "low"
        if delta_median <= P_HIGH_IMPACT_KBAR:
            return "moderate"
        return "high"
    if delta_median <= T_LOW_IMPACT_C:
        return "low"
    if delta_median <= T_HIGH_IMPACT_C:
        return "moderate"
    return "high"


def compute_impact_summary(pt_rows: pd.DataFrame, branch: BranchData, scope: str = "favored_subset") -> dict[str, Any]:
    """Compute across-model P-T impact metrics for a branch."""
    subset = pt_rows[(pt_rows["branch"] == branch.key.label) & (pt_rows["scope"] == scope)].copy()
    subset = subset[subset["model"].isin(branch.retained_models)]
    subset = subset[np.isfinite(subset["median"])]
    if subset.empty:
        return {
            "impact_scope": scope,
            "delta_median": np.nan,
            "delta_mean": np.nan,
            "delta_IQR_center": np.nan,
            "common_IQR_overlap": np.nan,
            "impact_level": "unknown",
            "pairwise_median_diffs_top_retained": "",
        }

    delta_median = float(subset["median"].max() - subset["median"].min())
    delta_mean = float(subset["mean"].max() - subset["mean"].min())
    delta_iqr_center = float(subset["q75"].max() - subset["q25"].min())
    iqr_span = float(subset["q75"].max() - subset["q25"].min())
    overlap = max(0.0, float(subset["q75"].min() - subset["q25"].max())) / max(1e-9, iqr_span)

    median_map = dict(zip(subset["model"], subset["median"]))
    pair_texts = []
    for a, b in itertools.combinations(branch.retained_models[:4], 2):
        if a in median_map and b in median_map:
            pair_texts.append(f"{model_abbrev(a, branch.key.target)}-{model_abbrev(b, branch.key.target)}:{abs(median_map[a] - median_map[b]):.3g}")

    return {
        "impact_scope": scope,
        "delta_median": delta_median,
        "delta_mean": delta_mean,
        "delta_IQR_center": delta_iqr_center,
        "common_IQR_overlap": float(overlap),
        "impact_level": impact_level(branch.key.target, delta_median),
        "pairwise_median_diffs_top_retained": "; ".join(pair_texts),
    }


def topk_rows(branch: BranchData, pt_rows: pd.DataFrame, scope: str = "favored_subset") -> list[dict[str, Any]]:
    """Audit how top-k retained models change support and median range."""
    rows: list[dict[str, Any]] = []
    subset = pt_rows[(pt_rows["branch"] == branch.key.label) & (pt_rows["scope"] == scope)].copy()
    median_map = dict(zip(subset["model"], subset["median"]))
    threshold = P_LOW_IMPACT_KBAR if branch.key.target == "P" else T_LOW_IMPACT_C
    unit = "kbar" if branch.key.target == "P" else "degC"
    props: pd.Series = branch.support.get("props", pd.Series(dtype=float))
    prev_min = np.nan
    prev_max = np.nan
    prev_delta = 0.0
    for k in range(1, min(4, len(branch.retained_models)) + 1):
        models_k = branch.retained_models[:k]
        medians = [median_map.get(model, np.nan) for model in models_k]
        medians = [float(value) for value in medians if np.isfinite(value)]
        if medians:
            min_median = float(np.min(medians))
            max_median = float(np.max(medians))
            delta = max_median - min_median
        else:
            min_median = np.nan
            max_median = np.nan
            delta = np.nan
        cumulative = float(sum(float(props.get(model, 0.0)) for model in models_k))
        expands = bool(k > 1 and np.isfinite(delta) and (delta - prev_delta) > threshold)
        new_tail = "none"
        kth_model = models_k[-1]
        kth_median = median_map.get(kth_model, np.nan)
        if k > 1 and np.isfinite(kth_median):
            if np.isfinite(prev_max) and kth_median > prev_max + threshold:
                new_tail = "high_pressure_tail" if branch.key.target == "P" else "high_temperature_tail"
            elif np.isfinite(prev_min) and kth_median < prev_min - threshold:
                new_tail = "low_pressure_tail" if branch.key.target == "P" else "low_temperature_tail"
        rows.append(
            {
                "branch": branch.key.label,
                "target": branch.key.target,
                "unit": unit,
                "phase_type": branch.key.phase_type,
                "model_type": branch.key.phase_label,
                "eruption_year": branch.key.year,
                "scope": scope,
                "k": k,
                "top_k_models": "; ".join(models_k),
                "top_k_model_abbrevs": "; ".join(model_abbrev(model, branch.key.target) for model in models_k),
                "cumulative_proportion_Ck": cumulative,
                "delta_median_k": delta,
                "min_model_median_k": min_median,
                "max_model_median_k": max_median,
                "central_range_k": f"{min_median:.3g}-{max_median:.3g}" if np.isfinite(min_median) else "",
                "expands_median_range_more_than_threshold": expands,
                "expansion_threshold": threshold,
                "new_tail_direction": new_tail,
            }
        )
        if np.isfinite(min_median):
            prev_min = min_median
            prev_max = max_median
            prev_delta = delta
    return rows


def composition_variables_for_branch(key: BranchKey) -> list[tuple[str, str]]:
    """Return preselected composition variables for a branch."""
    variables = [
        ("cpx_Al2O3", "cpx Al2O3"),
        ("cpx_Na2O", "cpx Na2O"),
        ("cpx_CaO", "cpx CaO"),
        ("cpx_MgNumber", "cpx Mg#"),
        ("cpx_SiO2", "cpx SiO2"),
        ("cpx_MgO", "cpx MgO"),
    ]
    if key.phase_type == "cpx_liq":
        variables.extend(
            [
                ("liq_SiO2", "liquid SiO2"),
                ("liq_TotalAlkali", "liquid Na2O+K2O"),
                ("liq_MgNumber", "liquid Mg#"),
                ("liq_MgO", "liquid MgO"),
                ("liq_H2O", "liquid H2O"),
            ]
        )
    return variables


def classify_composition_separation(max_sep: float) -> str:
    """Classify simple composition separation strength."""
    if not np.isfinite(max_sep):
        return "unknown"
    if max_sep < 0.5:
        return "weak"
    if max_sep < 0.8:
        return "moderate"
    if max_sep < 1.0:
        return "strong"
    return "very strong"


def compute_composition_metrics(branch: BranchData) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Compute group summaries and pairwise separation indices."""
    rows: list[dict[str, Any]] = []
    df = branch.frame.copy()
    variables = [(col, label) for col, label in composition_variables_for_branch(branch.key) if col in df.columns and df[col].notna().any()]
    retained = branch.retained_models
    group_stats: dict[tuple[str, str], dict[str, Any]] = {}

    for model in retained:
        group = df[df["Selected_model"].astype(str).eq(model)]
        for col, label in variables:
            values = finite_values(group[col])
            stats = summarize_numeric_values(values, use_tukey=False)
            group_stats[(model, col)] = stats
            rows.append(
                {
                    "branch": branch.key.label,
                    "target": branch.key.target,
                    "phase_type": branch.key.phase_type,
                    "model_type": branch.key.phase_label,
                    "eruption_year": branch.key.year,
                    "row_type": "group_summary",
                    "model": model,
                    "model_abbrev": model_abbrev(model, branch.key.target),
                    "model_a": "",
                    "model_b": "",
                    "model_pair_abbrev": "",
                    "variable": col,
                    "variable_label": label,
                    "n": stats["n_valid"],
                    "median": stats["median"],
                    "q25": stats["q25"],
                    "q75": stats["q75"],
                    "iqr": stats["iqr"],
                    "iqr_all": np.nan,
                    "sep": np.nan,
                }
            )

    sep_rows: list[dict[str, Any]] = []
    for model_a, model_b in itertools.combinations(retained, 2):
        for col, label in variables:
            all_values = finite_values(df[col])
            if all_values.size == 0:
                iqr_all = np.nan
            else:
                q1, q3 = np.nanpercentile(all_values, [25, 75])
                iqr_all = float(q3 - q1)
            med_a = group_stats.get((model_a, col), {}).get("median", np.nan)
            med_b = group_stats.get((model_b, col), {}).get("median", np.nan)
            sep = np.nan
            if np.isfinite(med_a) and np.isfinite(med_b) and np.isfinite(iqr_all) and iqr_all > 0:
                sep = float(abs(med_a - med_b) / iqr_all)
            sep_rows.append(
                {
                    "branch": branch.key.label,
                    "target": branch.key.target,
                    "phase_type": branch.key.phase_type,
                    "model_type": branch.key.phase_label,
                    "eruption_year": branch.key.year,
                    "row_type": "pair_sep",
                    "model": "",
                    "model_abbrev": "",
                    "model_a": model_a,
                    "model_b": model_b,
                    "model_pair_abbrev": f"{model_abbrev(model_a, branch.key.target)}-{model_abbrev(model_b, branch.key.target)}",
                    "variable": col,
                    "variable_label": label,
                    "n": np.nan,
                    "median": np.nan,
                    "q25": np.nan,
                    "q75": np.nan,
                    "iqr": np.nan,
                    "iqr_all": iqr_all,
                    "sep": sep,
                }
            )

    sep_values = [row for row in sep_rows if np.isfinite(row["sep"])]
    sep_values_sorted = sorted(sep_values, key=lambda row: row["sep"], reverse=True)
    best = sep_values_sorted[0] if sep_values_sorted else {}
    second = sep_values_sorted[1] if len(sep_values_sorted) > 1 else {}
    max_sep_by_variable: dict[str, float] = {}
    for row in sep_values:
        variable_key = str(row.get("variable_label") or row.get("variable"))
        current = max_sep_by_variable.get(variable_key, -np.inf)
        max_sep_by_variable[variable_key] = max(current, float(row["sep"]))

    summary = {
        "composition_sep_max": float(best.get("sep", np.nan)),
        "composition_sep_second": float(second.get("sep", np.nan)),
        "best_variable": best.get("variable_label", ""),
        "best_model_pair": best.get("model_pair_abbrev", ""),
        "n_variables_sep_ge_0p5": int(sum(1 for value in max_sep_by_variable.values() if value >= 0.5)),
        "n_variables_sep_ge_0p8": int(sum(1 for value in max_sep_by_variable.values() if value >= 0.8)),
        "n_variables_sep_ge_1p0": int(sum(1 for value in max_sep_by_variable.values() if value >= 1.0)),
    }
    summary["composition_separation_level"] = classify_composition_separation(summary["composition_sep_max"])
    summary["multi_variable_support"] = bool(summary["n_variables_sep_ge_0p5"] >= 2)

    rows.extend(sep_rows)
    rows.append(
        {
            "branch": branch.key.label,
            "target": branch.key.target,
            "phase_type": branch.key.phase_type,
            "model_type": branch.key.phase_label,
            "eruption_year": branch.key.year,
            "row_type": "branch_summary",
            "model": "",
            "model_abbrev": "",
            "model_a": "",
            "model_b": "",
            "model_pair_abbrev": summary["best_model_pair"],
            "variable": "",
            "variable_label": summary["best_variable"],
            "n": np.nan,
            "median": np.nan,
            "q25": np.nan,
            "q75": np.nan,
            "iqr": np.nan,
            "iqr_all": np.nan,
            "sep": summary["composition_sep_max"],
            **summary,
        }
    )
    return rows, summary


def scheme1_label(support: dict[str, Any]) -> str:
    """Apply proportion-only classification."""
    p1 = support.get("p1", 0.0)
    gap12 = support.get("gap12", 0.0)
    c2 = support.get("C2", 0.0)
    k80 = support.get("K80", np.nan)
    n_eff = support.get("N_eff", np.nan)
    if p1 >= DOMINANT_P1 and gap12 >= GAP_DOMINANT:
        return "dominant"
    if NEAR_DOMINANT_P1 <= p1 < DOMINANT_P1 and gap12 >= GAP_DOMINANT:
        return "near-dominant"
    if p1 < HIGHLY_DISPERSED_P1 and np.isfinite(k80) and k80 >= 3:
        return "highly dispersed"
    if p1 < NEAR_DOMINANT_P1 and c2 >= TWO_MODEL_COVERAGE:
        return "two-model competition"
    if (np.isfinite(k80) and k80 >= 3) or (np.isfinite(n_eff) and n_eff >= HIGH_DISPERSION_NEFF):
        return "dispersed"
    return "mixed support"


def scheme2_label(support: dict[str, Any], impact: dict[str, Any]) -> str:
    """Apply proportion plus P-T impact classification."""
    scheme1 = scheme1_label(support)
    impact_level_value = impact.get("impact_level", "unknown")
    c2 = support.get("C2", 0.0)
    k80 = support.get("K80", np.nan)
    if scheme1 == "dominant":
        return "dominant"
    if impact_level_value == "low":
        return "low-impact mixed selection"
    if c2 >= TWO_MODEL_COVERAGE and impact_level_value == "high":
        return "high-impact two-model competition"
    if c2 >= TWO_MODEL_COVERAGE and impact_level_value == "moderate":
        return "moderate-impact two-model competition"
    if np.isfinite(k80) and k80 >= 3 and impact_level_value == "high":
        return "high-impact dispersed/model-dependent"
    if np.isfinite(k80) and k80 >= 3 and impact_level_value == "moderate":
        return "moderate-impact dispersed"
    return f"mixed support with {impact_level_value} impact"


def scheme3_label(support: dict[str, Any], impact: dict[str, Any], composition: dict[str, Any]) -> str:
    """Apply proportion, P-T impact, and composition classification."""
    p1 = support.get("p1", 0.0)
    gap12 = support.get("gap12", 0.0)
    k80 = support.get("K80", np.nan)
    n_eff = support.get("N_eff", np.nan)
    impact_level_value = impact.get("impact_level", "unknown")
    comp_level = composition.get("composition_separation_level", "unknown")
    dispersed_support = (np.isfinite(k80) and k80 >= 3) or (np.isfinite(n_eff) and n_eff >= HIGH_DISPERSION_NEFF)

    if p1 >= DOMINANT_P1 and gap12 >= GAP_DOMINANT:
        return "resolved dominance"
    if NEAR_DOMINANT_P1 <= p1 < DOMINANT_P1 and gap12 >= GAP_DOMINANT:
        return "near-dominance"
    if impact_level_value == "low":
        return "low-impact mixed selection"
    if dispersed_support and impact_level_value in {"moderate", "high"}:
        return "dispersed/model-dependent"
    if impact_level_value in {"moderate", "high"} and comp_level in {"strong", "very strong"}:
        return "composition-dependent"
    if impact_level_value in {"moderate", "high"} and comp_level == "moderate":
        return "weakly composition-associated"
    if impact_level_value in {"moderate", "high"} and comp_level == "weak":
        return "ambiguous"
    return "mixed selection"


def recommended_action(label: str) -> str:
    """Return a manuscript action phrase from the final label."""
    if label in {"resolved dominance", "near-dominance"}:
        return "report the dominant model as the main result"
    if label == "low-impact mixed selection":
        return "report the overall favored-model distribution and note that model choice has limited impact"
    if label == "composition-dependent":
        return "report the overall favored-model distribution and only link it to geology if textures or zones support it"
    if label == "weakly composition-associated":
        return "report the overall favored-model distribution with a weak composition qualifier"
    if label == "ambiguous":
        return "use additional criteria to choose any representative model"
    if label == "dispersed/model-dependent":
        return "treat the result as a model-dependent scenario"
    return "report the mixed selection with the stated qualifiers"


def manuscript_sentence(branch: BranchData, impact: dict[str, Any], composition: dict[str, Any], label: str) -> str:
    """Build one manuscript-ready interpretation sentence."""
    support = branch.support
    lead_model = branch.retained_models[0] if branch.retained_models else "no retained model"
    lead = model_abbrev(lead_model, branch.key.target)
    unit = "kbar" if branch.key.target == "P" else "degC"
    support_phrase = f"{lead} accounts for {support.get('p1', np.nan) * 100:.1f}% of samples with gap12={support.get('gap12', np.nan):.2f}"
    impact_phrase = f"retained favored-subset medians differ by {impact.get('delta_median', np.nan):.2f} {unit} ({impact.get('impact_level', 'unknown')} impact)"
    comp_phrase = (
        f"composition_sep_max={composition.get('composition_sep_max', np.nan):.2f} "
        f"({composition.get('composition_separation_level', 'unknown')}, best {composition.get('best_variable', '')} "
        f"for {composition.get('best_model_pair', '')})"
    )
    return (
        f"{branch.key.label} is classified as {label} because {support_phrase}, "
        f"{impact_phrase}, and {comp_phrase}. Therefore, we {recommended_action(label)}."
    )


def sanity_note(branch: BranchData, final_label: str, impact: dict[str, Any], composition: dict[str, Any]) -> str:
    """Compare mechanical output with expected sanity-check behavior."""
    expected = EXPECTED_BEHAVIOR.get((branch.key.target, branch.key.phase_type, branch.key.year), "")
    if not expected:
        return ""
    metrics = (
        f"Observed {final_label}; p1={branch.support.get('p1', np.nan):.2f}, "
        f"K80={branch.support.get('K80', np.nan):.0f}, N_eff={branch.support.get('N_eff', np.nan):.2f}, "
        f"delta_median={impact.get('delta_median', np.nan):.2f}, "
        f"composition={composition.get('composition_separation_level', 'unknown')}."
    )
    return f"{expected} {metrics}"


def case_metric_row(
    branch: BranchData,
    impact: dict[str, Any],
    composition: dict[str, Any],
    scheme1: str,
    scheme2: str,
    scheme3: str,
    sentence: str,
    sanity: str,
) -> dict[str, Any]:
    """Combine all branch-level metrics into one row."""
    support = branch.support
    summary_rows = branch.summary_file_rows
    leading_model = branch.retained_models[0] if branch.retained_models else ""
    existing_tukey_note = "merapi_fig8_9_model_summary present" if not summary_rows.empty else "not found"
    return {
        "branch": branch.key.label,
        "target": branch.key.target,
        "unit": "kbar" if branch.key.target == "P" else "degC",
        "phase_type": branch.key.phase_type,
        "model_type": branch.key.phase_label,
        "eruption_year": branch.key.year,
        "report_file": str(branch.report_path),
        "N": support.get("N", 0),
        "K": support.get("K", 0),
        "retained_models": "; ".join(branch.retained_models),
        "retained_model_abbrevs": "; ".join(model_abbrev(model, branch.key.target) for model in branch.retained_models),
        "leading_model": leading_model,
        "leading_model_abbrev": model_abbrev(leading_model, branch.key.target) if leading_model else "",
        "p1": support.get("p1", np.nan),
        "p2": support.get("p2", np.nan),
        "p3": support.get("p3", np.nan),
        "p4": support.get("p4", np.nan),
        "gap12": support.get("gap12", np.nan),
        "C2": support.get("C2", np.nan),
        "C3": support.get("C3", np.nan),
        "C4": support.get("C4", np.nan),
        "K80": support.get("K80", np.nan),
        "K90": support.get("K90", np.nan),
        "N_eff": support.get("N_eff", np.nan),
        "H_norm": support.get("H_norm", np.nan),
        "other_prop": support.get("other_prop", np.nan),
        "retained_support_sum": support.get("retained_support_sum", np.nan),
        "impact_scope": impact.get("impact_scope", ""),
        "delta_median": impact.get("delta_median", np.nan),
        "delta_mean": impact.get("delta_mean", np.nan),
        "delta_IQR_center": impact.get("delta_IQR_center", np.nan),
        "common_IQR_overlap": impact.get("common_IQR_overlap", np.nan),
        "impact_level": impact.get("impact_level", "unknown"),
        "pairwise_median_diffs_top_retained": impact.get("pairwise_median_diffs_top_retained", ""),
        **composition,
        "scheme1_proportion_only": scheme1,
        "scheme2_proportion_pt_impact": scheme2,
        "scheme3_proportion_pt_composition": scheme3,
        "recommended_action": recommended_action(scheme3),
        "manuscript_sentence": sentence,
        "sanity_check_note": sanity,
        "pt_value_filter": "1.5 IQR Tukey fence recomputed" if USE_TUKEY_FOR_PT else "all finite values",
        "existing_summary_tukey_note": existing_tukey_note,
    }


def classification_row(branch: BranchData, impact: dict[str, Any], composition: dict[str, Any], scheme1: str, scheme2: str, scheme3: str, sentence: str, sanity: str) -> dict[str, Any]:
    """Return a compact branch classification row."""
    return {
        "branch": branch.key.label,
        "target": branch.key.target,
        "phase_type": branch.key.phase_type,
        "model_type": branch.key.phase_label,
        "eruption_year": branch.key.year,
        "scheme1_proportion_only": scheme1,
        "scheme2_proportion_pt_impact": scheme2,
        "scheme3_proportion_pt_composition": scheme3,
        "impact_level": impact.get("impact_level", "unknown"),
        "composition_separation_level": composition.get("composition_separation_level", "unknown"),
        "recommended_action": recommended_action(scheme3),
        "manuscript_sentence": sentence,
        "sanity_check_note": sanity,
    }


def write_csvs(output_dir: Path, tables: dict[str, pd.DataFrame]) -> None:
    """Write required CSV outputs."""
    output_dir.mkdir(parents=True, exist_ok=True)
    tables["case_metrics"].to_csv(output_dir / "merapi_selection_case_metrics.csv", index=False)
    tables["classification"].to_csv(output_dir / "merapi_selection_case_classification.csv", index=False)
    tables["topk"].to_csv(output_dir / "merapi_topk_scenario_audit.csv", index=False)
    tables["composition"].to_csv(output_dir / "merapi_composition_separation_metrics.csv", index=False)
    tables["model_distribution"].to_csv(output_dir / "merapi_model_distribution_metrics.csv", index=False)


def autosize_openpyxl_columns(writer: pd.ExcelWriter, sheet_name: str, df: pd.DataFrame) -> None:
    """Apply compact Excel formatting with openpyxl."""
    from openpyxl.styles import Font, PatternFill

    worksheet = writer.sheets[sheet_name]
    worksheet.freeze_panes = "A2"
    for idx, col in enumerate(df.columns, start=1):
        values = [str(col)]
        if not df.empty:
            values.extend(str(value) for value in df[col].head(200).fillna("").tolist())
        width = min(max(max(len(value) for value in values) + 2, 10), 42)
        worksheet.column_dimensions[worksheet.cell(row=1, column=idx).column_letter].width = width
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")


def write_xlsx(output_dir: Path, tables: dict[str, pd.DataFrame]) -> Path:
    """Write the required XLSX workbook."""
    xlsx_path = output_dir / "merapi_selection_case_metrics.xlsx"
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        sheet_map = {
            "case_metrics": "case_metrics",
            "classification": "classification",
            "topk": "topk_scenario",
            "composition": "composition_sep",
            "model_distribution": "model_distribution",
            "pt_distribution": "pt_distribution",
            "input_diagnostics": "input_diagnostics",
            "warnings": "warnings",
        }
        for table_key, sheet_name in sheet_map.items():
            df = tables.get(table_key, pd.DataFrame())
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            autosize_openpyxl_columns(writer, sheet_name, df)
    return xlsx_path


def pil_font(size: int, *, bold: bool = False):
    """Return a readable PIL font with a safe default fallback."""
    from PIL import ImageFont

    candidates = [
        "arialbd.ttf" if bold else "arial.ttf",
        "segoeuib.ttf" if bold else "segoeui.ttf",
        "calibrib.ttf" if bold else "calibri.ttf",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size=size)
        except Exception:
            continue
    return ImageFont.load_default()


def hex_to_rgb(color: str) -> tuple[int, int, int]:
    """Convert a #RRGGBB color to an RGB tuple."""
    color = color.strip().lstrip("#")
    return tuple(int(color[i : i + 2], 16) for i in (0, 2, 4))


def lerp_color(a: str, b: str, t: float) -> tuple[int, int, int]:
    """Blend two hex colors."""
    t = max(0.0, min(1.0, float(t)))
    ar, ag, ab = hex_to_rgb(a)
    br, bg, bb = hex_to_rgb(b)
    return (int(ar + (br - ar) * t), int(ag + (bg - ag) * t), int(ab + (bb - ab) * t))


def save_pil_chart(image: Any, output_dir: Path, stem: str) -> tuple[Path, Path]:
    """Save a Pillow chart image as PNG and PDF."""
    output_dir.mkdir(parents=True, exist_ok=True)
    png = output_dir / f"{stem}.png"
    pdf = output_dir / f"{stem}.pdf"
    image.save(png)
    image.convert("RGB").save(pdf, "PDF", resolution=180.0)
    return png, pdf


def make_canvas(width: int, height: int, title: str):
    """Create a white chart canvas and draw its title."""
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((40, 24), title, fill="#111827", font=pil_font(28, bold=True))
    return image, draw


def panel_boxes(width: int, height: int, *, top: int = 84, rows: int = 4, cols: int = 2) -> list[tuple[int, int, int, int]]:
    """Return evenly spaced panel boxes."""
    margin_x = 42
    margin_bottom = 38
    gap_x = 34
    gap_y = 42
    panel_w = (width - 2 * margin_x - (cols - 1) * gap_x) // cols
    panel_h = (height - top - margin_bottom - (rows - 1) * gap_y) // rows
    boxes = []
    for r in range(rows):
        for c in range(cols):
            x0 = margin_x + c * (panel_w + gap_x)
            y0 = top + r * (panel_h + gap_y)
            boxes.append((x0, y0, x0 + panel_w, y0 + panel_h))
    return boxes


def draw_panel_frame(draw: Any, box: tuple[int, int, int, int], title: str) -> tuple[int, int, int, int]:
    """Draw a panel border and title, returning the inner plot area."""
    x0, y0, x1, y1 = box
    draw.rectangle(box, outline="#D1D5DB", width=1)
    draw.text((x0 + 12, y0 + 8), title, fill="#111827", font=pil_font(16, bold=True))
    return (x0 + 54, y0 + 42, x1 - 18, y1 - 42)


def draw_axes(draw: Any, plot: tuple[int, int, int, int], y_min: float, y_max: float, *, ticks: int = 4, suffix: str = "") -> None:
    """Draw simple x-y axes with y tick labels."""
    x0, y0, x1, y1 = plot
    draw.line((x0, y0, x0, y1, x1, y1), fill="#374151", width=1)
    font = pil_font(10)
    for i in range(ticks + 1):
        frac = i / ticks
        value = y_min + (y_max - y_min) * frac
        y = y1 - int(frac * (y1 - y0))
        draw.line((x0 - 4, y, x0, y), fill="#374151", width=1)
        label = f"{value:.0%}" if suffix == "%" else f"{value:.1f}{suffix}"
        draw.text((x0 - 50, y - 7), label, fill="#4B5563", font=font)
        if i:
            draw.line((x0 + 1, y, x1, y), fill="#E5E7EB", width=1)


def y_to_pixel(value: float, y_min: float, y_max: float, plot: tuple[int, int, int, int]) -> int:
    """Map a y value to a pixel coordinate."""
    _x0, y0, _x1, y1 = plot
    if not np.isfinite(value) or y_max <= y_min:
        return y1
    frac = (value - y_min) / (y_max - y_min)
    return y1 - int(max(0.0, min(1.0, frac)) * (y1 - y0))


def plot_model_proportions(model_distribution: pd.DataFrame, output_dir: Path) -> tuple[Path, Path]:
    """Plot model proportion bars for all branches using Pillow."""
    grouped = model_distribution[model_distribution["row_type"] == "grouped_for_support"].copy()
    branches = [BranchKey(*b).label for b in BRANCHES]
    image, draw = make_canvas(1800, 2200, "Merapi favored-model support proportions")
    for box, branch_label in zip(panel_boxes(1800, 2200), branches):
        plot = draw_panel_frame(draw, box, branch_label)
        data = grouped[grouped["branch"] == branch_label].sort_values("favored_proportion", ascending=False)
        labels = data["model_abbrev"].tolist()
        values = data["favored_proportion"].to_numpy(dtype=float)
        y_max = max(0.75, min(1.0, float(np.nanmax(values)) + 0.12 if values.size else 0.75))
        draw_axes(draw, plot, 0.0, y_max, suffix="%")
        x0, y0, x1, y1 = plot
        if not values.size:
            draw.text((x0 + 20, y0 + 50), "No support data", fill="#6B7280", font=pil_font(13))
            continue
        bar_gap = 8
        bar_w = max(18, (x1 - x0 - bar_gap * (len(values) + 1)) // max(1, len(values)))
        min_y = y_to_pixel(MIN_PROP_RETAIN, 0.0, y_max, plot)
        draw.line((x0, min_y, x1, min_y), fill="#9CA3AF", width=1)
        for i, (label, value) in enumerate(zip(labels, values)):
            bx0 = x0 + bar_gap + i * (bar_w + bar_gap)
            bx1 = bx0 + bar_w
            by = y_to_pixel(float(value), 0.0, y_max, plot)
            color = "#4C78A8" if label != "Other" else "#9CA3AF"
            draw.rectangle((bx0, by, bx1, y1), fill=color, outline="#374151")
            draw.text((bx0, by - 16), f"{value:.0%}", fill="#111827", font=pil_font(10))
            draw.text((bx0, y1 + 8), label, fill="#111827", font=pil_font(10))
    return save_pil_chart(image, output_dir, "merapi_model_proportion_bars")


def plot_model_median_iqr(pt_distribution: pd.DataFrame, output_dir: Path) -> tuple[Path, Path]:
    """Plot model median plus IQR for retained models using Pillow."""
    data_all = pt_distribution[pt_distribution["scope"] == "favored_subset"].copy()
    branches = [BranchKey(*b).label for b in BRANCHES]
    image, draw = make_canvas(1800, 2200, "Retained model favored-subset P-T medians and IQRs")
    for box, branch_label in zip(panel_boxes(1800, 2200), branches):
        data = data_all[data_all["branch"] == branch_label].sort_values("model_abbrev")
        unit = data["unit"].iloc[0] if not data.empty else ""
        plot = draw_panel_frame(draw, box, f"{branch_label} ({unit})")
        if data.empty:
            draw.text((plot[0] + 20, plot[1] + 50), "No retained model data", fill="#6B7280", font=pil_font(13))
            continue
        y_values = np.concatenate([data["q25"].to_numpy(dtype=float), data["q75"].to_numpy(dtype=float)])
        y_values = y_values[np.isfinite(y_values)]
        y_min = float(np.nanmin(y_values))
        y_max = float(np.nanmax(y_values))
        pad = max((y_max - y_min) * 0.12, 0.5 if unit == "kbar" else 15.0)
        y_min -= pad
        y_max += pad
        draw_axes(draw, plot, y_min, y_max)
        x0, y0, x1, y1 = plot
        labels = data["model_abbrev"].tolist()
        n = len(labels)
        for i, row in enumerate(data.itertuples(index=False)):
            x = x0 + int((i + 0.5) * (x1 - x0) / max(1, n))
            med = float(row.median)
            q25 = float(row.q25)
            q75 = float(row.q75)
            y_med = y_to_pixel(med, y_min, y_max, plot)
            y_q25 = y_to_pixel(q25, y_min, y_max, plot)
            y_q75 = y_to_pixel(q75, y_min, y_max, plot)
            draw.line((x, y_q25, x, y_q75), fill="#4C78A8", width=4)
            draw.line((x - 7, y_q25, x + 7, y_q25), fill="#4C78A8", width=2)
            draw.line((x - 7, y_q75, x + 7, y_q75), fill="#4C78A8", width=2)
            draw.ellipse((x - 5, y_med - 5, x + 5, y_med + 5), fill="#111827")
            draw.text((x - 18, y_med - 22), f"{med:.1f}", fill="#111827", font=pil_font(10))
            draw.text((x - 18, y1 + 8), labels[i], fill="#111827", font=pil_font(10))
    return save_pil_chart(image, output_dir, "merapi_model_median_iqr")


def plot_topk(topk: pd.DataFrame, output_dir: Path) -> tuple[Path, Path]:
    """Plot cumulative support and median spread as top-k models are added using Pillow."""
    branches = [BranchKey(*b).label for b in BRANCHES]
    image, draw = make_canvas(1800, 2200, "Top-k support coverage and P-T median range")
    blue = "#4C78A8"
    orange = "#D55E00"
    for box, branch_label in zip(panel_boxes(1800, 2200), branches):
        plot = draw_panel_frame(draw, box, branch_label)
        data = topk[topk["branch"] == branch_label].sort_values("k")
        if data.empty:
            draw.text((plot[0] + 20, plot[1] + 50), "No retained models", fill="#6B7280", font=pil_font(13))
            continue
        x0, y0, x1, y1 = plot
        draw_axes(draw, plot, 0.0, 1.0, suffix="%")
        k = data["k"].to_numpy(dtype=float)
        c = data["cumulative_proportion_Ck"].to_numpy(dtype=float)
        delta = data["delta_median_k"].to_numpy(dtype=float)
        delta_max = max(float(np.nanmax(delta)) if np.isfinite(delta).any() else 1.0, float(data["expansion_threshold"].iloc[0]) * 1.4)
        points_c = []
        points_d = []
        for i, kval in enumerate(k):
            x = x0 + int((kval - 1) / max(1, (max(k) - 1)) * (x1 - x0)) if max(k) > 1 else (x0 + x1) // 2
            points_c.append((x, y_to_pixel(float(c[i]), 0.0, 1.0, plot)))
            points_d.append((x, y_to_pixel(float(delta[i]) / delta_max, 0.0, 1.0, plot)))
            draw.text((x - 3, y1 + 8), str(int(kval)), fill="#111827", font=pil_font(10))
        if len(points_c) > 1:
            draw.line(points_c, fill=blue, width=3)
            draw.line(points_d, fill=orange, width=3)
        for x, y in points_c:
            draw.ellipse((x - 5, y - 5, x + 5, y + 5), fill=blue)
        for x, y in points_d:
            draw.rectangle((x - 5, y - 5, x + 5, y + 5), fill=orange)
        thresh_norm = float(data["expansion_threshold"].iloc[0]) / delta_max
        thresh_y = y_to_pixel(thresh_norm, 0.0, 1.0, plot)
        draw.line((x0, thresh_y, x1, thresh_y), fill=orange, width=1)
        unit = data["unit"].iloc[0]
        draw.text((x0 + 8, y0 + 8), "blue=Ck", fill=blue, font=pil_font(11, bold=True))
        draw.text((x0 + 95, y0 + 8), f"orange=delta median, max scale {delta_max:.2g} {unit}", fill=orange, font=pil_font(11, bold=True))
    return save_pil_chart(image, output_dir, "merapi_topk_scenario")


def plot_composition_heatmaps(composition: pd.DataFrame, output_dir: Path) -> tuple[Path, Path]:
    """Plot simple separation-index heatmaps by branch using Pillow."""
    pair_rows = composition[composition["row_type"] == "pair_sep"].copy()
    branches = [BranchKey(*b).label for b in BRANCHES]
    image, draw = make_canvas(2400, 2600, "Pairwise composition separation by retained model group")
    for box, branch_label in zip(panel_boxes(2400, 2600, rows=4, cols=2), branches):
        x0, y0, x1, y1 = box
        draw.rectangle(box, outline="#D1D5DB", width=1)
        draw.text((x0 + 12, y0 + 8), branch_label, fill="#111827", font=pil_font(16, bold=True))
        data = pair_rows[pair_rows["branch"] == branch_label].copy()
        data = data[np.isfinite(data["sep"])]
        if data.empty:
            draw.text((x0 + 24, y0 + 70), "No pairwise composition data", fill="#6B7280", font=pil_font(13))
            continue
        pivot = data.pivot_table(index="variable_label", columns="model_pair_abbrev", values="sep", aggfunc="max")
        pivot = pivot.loc[pivot.max(axis=1).sort_values(ascending=False).index]
        max_rows = min(8, len(pivot.index))
        max_cols = min(6, len(pivot.columns))
        pivot = pivot.iloc[:max_rows, :max_cols]
        label_w = 168
        top_h = 64
        cell_w = max(48, (x1 - x0 - label_w - 16) // max_cols)
        cell_h = max(34, (y1 - y0 - top_h - 22) // max_rows)
        table_x = x0 + label_w
        table_y = y0 + top_h
        for j, col in enumerate(pivot.columns):
            draw.text((table_x + j * cell_w + 2, y0 + 36), str(col)[:14], fill="#111827", font=pil_font(9))
        for i, var in enumerate(pivot.index):
            draw.text((x0 + 10, table_y + i * cell_h + 9), str(var)[:24], fill="#111827", font=pil_font(10))
            for j, col in enumerate(pivot.columns):
                value = pivot.loc[var, col]
                t = min(1.0, float(value) / 1.2) if np.isfinite(value) else 0.0
                color = lerp_color("#F3F4F6", "#256D85", t)
                cx0 = table_x + j * cell_w
                cy0 = table_y + i * cell_h
                draw.rectangle((cx0, cy0, cx0 + cell_w - 2, cy0 + cell_h - 2), fill=color, outline="white")
                if np.isfinite(value):
                    text_color = "white" if t > 0.55 else "#111827"
                    draw.text((cx0 + 9, cy0 + 9), f"{value:.2f}", fill=text_color, font=pil_font(10))
    draw.text((40, 2560), "Cell values are abs(group median difference) / branch-wide IQR; darker cells are more separated.", fill="#374151", font=pil_font(14))
    return save_pil_chart(image, output_dir, "merapi_composition_separation_heatmaps")


def write_summary_markdown(output_dir: Path, case_metrics: pd.DataFrame, warnings_df: pd.DataFrame, figure_paths: list[tuple[Path, Path]]) -> Path:
    """Write the manuscript-facing summary markdown."""
    path = output_dir / "merapi_selection_case_summary.md"
    lines: list[str] = []
    lines.append("# Merapi Model-Selection Pattern Audit")
    lines.append("")
    lines.append(
        "This audit classifies AIMS4PT_cpx sample-level model-selection patterns. It is not a geological clustering analysis, "
        "does not identify reservoirs, and does not use PCA, clustering, or machine-learning classifiers. Composition-dependent "
        "selection is therefore a model-selection pattern unless it can be tied to independent textures, zones, or populations."
    )
    lines.append("")
    lines.append("## Method Notes")
    lines.append("")
    lines.append(f"- Retained models have favored proportion >= {MIN_PROP_RETAIN:.0%}; lower-proportion models are grouped as Other for support summaries.")
    lines.append("- P-T impact uses favored-subset retained-model predictions after a recomputed 1.5 IQR Tukey fence, matching the Merapi helper convention.")
    lines.append("- cpx-liquid composition variables use cached KD=0.28 pairing rows that already passed the report filter; pairing is not recalculated.")
    lines.append("- Composition separation is abs(median_a - median_b) / branch-wide IQR for predefined variables only.")
    lines.append("")
    lines.append("## Output Figures")
    lines.append("")
    for png, pdf in figure_paths:
        lines.append(f"- {png.name}; {pdf.name}")
    lines.append("")
    lines.append("## Branch Summaries")
    lines.append("")

    for _, row in case_metrics.iterrows():
        lines.append(f"### {row['branch']}")
        lines.append("")
        lines.append(
            f"- Retained models: {row['retained_model_abbrevs']} ({row['retained_models']})"
        )
        lines.append(
            f"- Support: N={int(row['N'])}, K={int(row['K'])}, p1={row['p1']:.3f}, gap12={row['gap12']:.3f}, "
            f"C2={row['C2']:.3f}, C3={row['C3']:.3f}, K80={row['K80']:.0f}, K90={row['K90']:.0f}, "
            f"N_eff={row['N_eff']:.2f}, H_norm={row['H_norm']:.2f}"
        )
        lines.append(
            f"- P-T impact: delta_median={row['delta_median']:.2f} {row['unit']} "
            f"({row['impact_level']} impact; scope={row['impact_scope']})"
        )
        lines.append(
            f"- Composition: composition_sep_max={row['composition_sep_max']:.2f} "
            f"({row['composition_separation_level']}), best variable={row['best_variable']}, "
            f"best pair={row['best_model_pair']}"
        )
        lines.append(
            f"- Classifications: Scheme 1={row['scheme1_proportion_only']}; "
            f"Scheme 2={row['scheme2_proportion_pt_impact']}; "
            f"Scheme 3={row['scheme3_proportion_pt_composition']}"
        )
        lines.append(f"- Recommended sentence: {row['manuscript_sentence']}")
        lines.append(f"- Sanity check: {row['sanity_check_note']}")
        lines.append("")

    if not warnings_df.empty:
        lines.append("## Diagnostic Warnings")
        lines.append("")
        for _, row in warnings_df.iterrows():
            branch = f"[{row['branch']}] " if row.get("branch") else ""
            file_text = f" ({Path(row['file']).name})" if row.get("file") else ""
            lines.append(f"- {branch}{row['warning']}{file_text}")
        lines.append("")

    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def build_tables(cache_dir: Path, use_tukey: bool, warnings_out: list[dict[str, str]]) -> tuple[dict[str, pd.DataFrame], list[BranchData]]:
    """Read branch data and compute all output tables."""
    summary_path = find_summary_file(cache_dir)
    if summary_path is None:
        warn(warnings_out, "merapi_fig8_9_model_summary.xlsx was not found; using report predictions only.")
    branches: list[BranchData] = []
    for target, phase_type, year in BRANCHES:
        key = BranchKey(target, phase_type, year)
        branches.append(read_branch_data(cache_dir, summary_path, key, warnings_out))

    model_distribution = pd.DataFrame([row for branch in branches for row in model_distribution_rows(branch)])
    pt_distribution = pd.DataFrame([row for branch in branches for row in compute_pt_distribution_rows(branch, use_tukey=use_tukey, warnings_out=warnings_out)])

    topk_all: list[dict[str, Any]] = []
    composition_all: list[dict[str, Any]] = []
    case_rows: list[dict[str, Any]] = []
    class_rows: list[dict[str, Any]] = []
    for branch in branches:
        impact = compute_impact_summary(pt_distribution, branch, scope="favored_subset")
        comp_rows, comp_summary = compute_composition_metrics(branch)
        composition_all.extend(comp_rows)
        topk_all.extend(topk_rows(branch, pt_distribution, scope="favored_subset"))
        s1 = scheme1_label(branch.support)
        s2 = scheme2_label(branch.support, impact)
        s3 = scheme3_label(branch.support, impact, comp_summary)
        sentence = manuscript_sentence(branch, impact, comp_summary, s3)
        sanity = sanity_note(branch, s3, impact, comp_summary)
        case_rows.append(case_metric_row(branch, impact, comp_summary, s1, s2, s3, sentence, sanity))
        class_rows.append(classification_row(branch, impact, comp_summary, s1, s2, s3, sentence, sanity))

    tables = {
        "case_metrics": pd.DataFrame(case_rows),
        "classification": pd.DataFrame(class_rows),
        "topk": pd.DataFrame(topk_all),
        "composition": pd.DataFrame(composition_all),
        "model_distribution": model_distribution,
        "pt_distribution": pt_distribution,
    }
    return tables, branches


def main() -> int:
    configure_stdout()
    args = parse_args()
    repo_root = args.repo_root.resolve()
    output_dir = args.output_dir if args.output_dir.is_absolute() else repo_root / args.output_dir
    warnings_out: list[dict[str, str]] = []
    use_tukey = not args.no_tukey
    globals()["USE_TUKEY_FOR_PT"] = use_tukey

    cache_dir = find_cache_dir(repo_root, args.timestamp, args.cache_dir, warnings_out)
    print(f"Repository root: {repo_root}")
    print(f"Merapi cache directory: {cache_dir}")
    print(f"Output directory: {output_dir}")
    print(f"P-T value filter: {'1.5 IQR Tukey fence' if use_tukey else 'all finite values'}")

    input_files = detect_input_files(cache_dir)
    input_diagnostics = print_input_diagnostics(input_files)
    tables, _branches = build_tables(cache_dir, use_tukey, warnings_out)
    tables["input_diagnostics"] = input_diagnostics
    tables["warnings"] = pd.DataFrame(warnings_out)

    write_csvs(output_dir, tables)
    xlsx_path = write_xlsx(output_dir, tables)
    figure_paths = [
        plot_model_proportions(tables["model_distribution"], output_dir),
        plot_model_median_iqr(tables["pt_distribution"], output_dir),
        plot_topk(tables["topk"], output_dir),
        plot_composition_heatmaps(tables["composition"], output_dir),
    ]
    summary_path = write_summary_markdown(output_dir, tables["case_metrics"], tables["warnings"], figure_paths)

    output_paths = [
        output_dir / "merapi_selection_case_metrics.csv",
        xlsx_path,
        output_dir / "merapi_selection_case_classification.csv",
        output_dir / "merapi_topk_scenario_audit.csv",
        output_dir / "merapi_composition_separation_metrics.csv",
        output_dir / "merapi_model_distribution_metrics.csv",
        summary_path,
    ]
    output_paths.extend(path for pair in figure_paths for path in pair)

    print("\nWrote Merapi selection audit outputs")
    print("-----------------------------------")
    for path in output_paths:
        print(path)

    key_cols = [
        "branch",
        "retained_model_abbrevs",
        "p1",
        "K80",
        "N_eff",
        "delta_median",
        "impact_level",
        "composition_separation_level",
        "scheme3_proportion_pt_composition",
    ]
    print("\nClassification preview")
    print("----------------------")
    print(tables["case_metrics"][key_cols].to_string(index=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
