"""Regression tests for the version-two Excel report."""

from __future__ import annotations

import io
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from pandas.testing import assert_frame_equal

from aims4pt.reporting import excel
from aims4pt.reporting import excel_v2
from aims4pt.reporting.excel import ReportPayload
from aims4pt.reporting.model_abbreviations import model_abbreviation_lookup


MODEL_A = "Petrelli et al. (2020) (cpx_only)"
MODEL_B = "Petrelli et al. (2020) (cpx_liq)"
MODEL_C = "Doe et al. (2019)"
MODEL_D = "Roe et al. (2018)"
MODEL_NAMES = [MODEL_A, MODEL_B, MODEL_C, MODEL_D]
PU08_A = "Putirka, 2008 eq32d_T; eq32a_P (cpx_only)"
PU08_B = "Putirka, 2008 eq32d_T; eq32b_P (cpx_only)"


@dataclass
class FakeModel:
    model_name: str
    T_P: str

    def __post_init__(self) -> None:
        self.X_cpx_training = pd.DataFrame({"SiO2": [49.0, 50.0, 51.0]})
        self.X_cpx_all = pd.DataFrame({"SiO2": np.linspace(45.0, 54.0, 123)})
        self.y_min = 0.54 if self.T_P == "P" else 899.6
        self.y_max = 12.34 if self.T_P == "P" else 1400.6
        self.rock_types = ["Basalt", "Basaltic andesite", "Dacite"]
        self.uncertainty = None

    def export_composition_range(self, *_args, **_kwargs) -> str:
        return (
            "SiO2: 45.0–54.0; TiO2: 0.2–2.1; Al2O3: 1.5–10.2; "
            "FeO: 2.0–12.0; MgO: 5.0–19.0; CaO: 8.0–24.0"
        )


class FakeWorkflow:
    def __init__(self, target: str) -> None:
        self.T_P = target
        if target == "P":
            self.prediction_df = pd.DataFrame(
                {
                    MODEL_A: np.linspace(1.24, 9.86, 20),
                    MODEL_B: np.linspace(2.16, 10.74, 20),
                    MODEL_C: np.linspace(3.04, 11.95, 20),
                    MODEL_D: np.full(20, 6.25),
                }
            )
            self.calculated_deviation_df = pd.DataFrame(
                {
                    MODEL_A: np.linspace(0.14, 1.04, 20),
                    MODEL_B: np.linspace(0.25, 1.45, 20),
                    MODEL_C: np.full(20, 0.64),
                    MODEL_D: np.full(20, 0.84),
                }
            )
        else:
            self.prediction_df = pd.DataFrame(
                {
                    MODEL_A: np.linspace(1000.4, 1219.6, 20),
                    MODEL_B: np.linspace(1012.6, 1233.4, 20),
                    MODEL_C: np.linspace(990.4, 1177.7, 20),
                    MODEL_D: np.full(20, 1105.4),
                }
            )
            self.calculated_deviation_df = pd.DataFrame(
                {
                    MODEL_A: np.linspace(18.4, 28.8, 20),
                    MODEL_B: np.linspace(22.2, 32.6, 20),
                    MODEL_C: np.full(20, 23.6),
                    MODEL_D: np.full(20, 27.6),
                }
            )
        self.ood_mask_df = pd.DataFrame(
            {
                name: [index % (model_index + 3) == 0 for index in range(20)]
                for model_index, name in enumerate(MODEL_NAMES)
            }
        )
        self.failure_reason_df = pd.DataFrame(
            {"Reason": ["" if index < 19 else "No eligible model" for index in range(20)]}
        )
        self._selected = pd.Series(
            [MODEL_A] * 14 + [MODEL_B] * 4 + [MODEL_C, None],
            name="Selected_model",
        )

    def get_best_model_series(self) -> pd.Series:
        return self._selected.copy()


def make_models(target: str) -> list[FakeModel]:
    return [FakeModel(name, target) for name in MODEL_NAMES]


def make_composition_data(n_rows: int = 20, *, include_liquid: bool = True) -> pd.DataFrame:
    phase = np.linspace(0.0, 2.0 * np.pi, n_rows, endpoint=False)
    data = pd.DataFrame(
        {
            "CaO_cpx": 16.0 + 3.0 * np.sin(phase),
            "MgO_cpx": 14.0 + 2.0 * np.cos(phase),
            "FeO_cpx": 6.0 + 1.2 * np.sin(phase + 0.4),
            "Na2O_cpx": 0.45 + 0.18 * np.cos(phase),
            "Al2O3_cpx": 5.0 + 1.8 * np.sin(phase - 0.3),
        }
    )
    if include_liquid:
        data = data.assign(
            SiO2_liq=57.0 + 8.0 * np.sin(phase),
            Na2O_liq=3.2 + 0.7 * np.cos(phase),
            K2O_liq=1.8 + 0.6 * np.sin(phase + 0.5),
            MgO_liq=4.5 + 1.4 * np.cos(phase - 0.2),
            FeO_liq=7.5 + 1.0 * np.sin(phase + 0.6),
        )
    return data


@pytest.fixture(scope="module")
def independent_data() -> pd.DataFrame:
    return make_composition_data(36)


@pytest.fixture(scope="module")
def pressure_payload(independent_data: pd.DataFrame) -> ReportPayload:
    return excel_v2.build_report_payload(
        FakeWorkflow("P"),
        make_models("P"),
        make_composition_data(),
        independent_data=independent_data,
        image_dpi=100,
    )


def test_packaged_reference_data_and_kde_cache_are_valid() -> None:
    reference_df, kde_cache = excel_v2._load_default_reference_context()

    assert len(reference_df) > 100
    for spec in excel_v2.COMPOSITION_PANEL_SPECS:
        key = spec["key"]
        assert kde_cache[f"{key}__xx"].shape == (120, 120)
        assert kde_cache[f"{key}__density"].shape == (120, 120)


def test_legacy_entry_and_compact_web_tables_are_preserved(
    monkeypatch: pytest.MonkeyPatch,
    pressure_payload: ReportPayload,
) -> None:
    monkeypatch.setattr(excel, "_build_violin_png", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        excel, "_build_model_selection_png", lambda *_args, **_kwargs: None
    )
    monkeypatch.setattr(
        excel, "_build_deviation_violin_png", lambda *_args, **_kwargs: None
    )
    legacy = excel.build_report_payload_legacy(
        FakeWorkflow("P"), make_models("P"), make_composition_data()
    )

    assert_frame_equal(pressure_payload.results_sheet, legacy.results_sheet)
    assert_frame_equal(pressure_payload.ranking_details_df, legacy.ranking_details_df)
    assert_frame_equal(pressure_payload.model_summary_df, legacy.model_summary_df)
    assert_frame_equal(pressure_payload.model_votes_df, legacy.model_votes_df)
    assert "Composition_range (wt%)" not in legacy.model_summary_df
    assert legacy.excel_model_summary_df is None


def test_full_summary_precision_abbreviations_and_vote_denominator(
    pressure_payload: ReportPayload,
) -> None:
    summary = pressure_payload.excel_model_summary_df
    votes = pressure_payload.excel_model_votes_df
    assert summary is not None
    assert votes is not None
    assert list(summary.columns[:6]) == [
        "Model_name",
        "Model_abbreviation",
        "Num_calibration_experiments",
        "Composition_range (wt%)",
        "P_range (kbar)",
        "Supported_TAS_rock_types",
    ]
    assert summary.loc[0, "Model_abbreviation"] == "Pet20"
    assert summary.loc[0, "P_range (kbar)"] == "0.5–12.3"
    assert summary.loc[0, "P_min (kbar)"] == pytest.approx(1.2)
    assert summary.loc[0, "Mean_calculated_deviation"] == pytest.approx(0.6)
    assert summary.loc[0, "OOD_ratio (%)"] == pytest.approx(35.0)
    assert pressure_payload.model_selection_image_dpi == 150
    assert pressure_payload.prediction_image_dpi == 150
    assert pressure_payload.deviation_image_dpi == 100

    no_model = votes.loc[votes["Model_name"] == "No selected model"].iloc[0]
    total = votes.loc[votes["Model_name"] == "Total"].iloc[0]
    assert no_model["Model_abbreviation"] == "—"
    assert no_model["Proportion"] == pytest.approx(0.05)
    assert total["Model_abbreviation"] == "—"
    assert total["Votes count"] == 20
    assert total["Proportion"] == pytest.approx(1.0)


def test_temperature_summary_uses_integer_precision(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        excel_v2, "_build_model_selection_png", lambda *_args, **_kwargs: None
    )
    monkeypatch.setattr(
        excel_v2, "_build_prediction_distribution_png", lambda *_args, **_kwargs: None
    )
    monkeypatch.setattr(
        excel_v2, "_build_deviation_violin_png", lambda *_args, **_kwargs: None
    )
    payload = excel_v2.build_report_payload(
        FakeWorkflow("T"),
        make_models("T"),
        make_composition_data(),
        independent_data=make_composition_data(24),
    )
    summary = payload.excel_model_summary_df
    assert summary is not None
    assert summary.loc[0, "T_range (C)"] == "900–1401"
    assert summary.loc[0, "T_min (C)"] == pytest.approx(1000.0)
    assert summary.loc[2, "Mean_calculated_deviation"] == pytest.approx(24.0)


def test_prediction_plot_uses_raw_overall_all_and_favored_values(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workflow = FakeWorkflow("P")
    records: list[tuple[np.ndarray, float, dict[str, object]]] = []
    captured_figures = []

    def record_distribution(_axis, values, position, **kwargs):
        records.append((excel_v2._finite_values(values), position, kwargs))
        return True

    def capture_figure(fig, dpi=200):
        captured_figures.append(fig)
        return b"png"

    monkeypatch.setattr(excel_v2, "_draw_distribution", record_distribution)
    monkeypatch.setattr(excel_v2, "_figure_to_png", capture_figure)
    excel_v2._build_prediction_distribution_png(
        workflow, make_models("P"), "P", 0.05, 100
    )

    assert len(records) == 1 + 2 * len(MODEL_NAMES)
    expected_overall = np.concatenate(
        [
            workflow.prediction_df.loc[:13, MODEL_A].to_numpy(),
            workflow.prediction_df.loc[14:17, MODEL_B].to_numpy(),
            workflow.prediction_df.loc[[18], MODEL_C].to_numpy(),
        ]
    )
    np.testing.assert_allclose(records[0][0], expected_overall)
    np.testing.assert_allclose(records[1][0], workflow.prediction_df[MODEL_A])
    np.testing.assert_allclose(
        records[2][0], workflow.prediction_df.loc[:13, MODEL_A]
    )
    assert records[6][0].size == 1
    assert records[8][0].size == 0
    assert records[1][2]["linestyle"] == "--"
    assert records[3][2]["linestyle"] == "--"
    assert records[5][2]["linestyle"] == "-"
    assert records[1][2]["linewidth"] == pytest.approx(2.0)
    assert records[2][2]["linewidth"] == pytest.approx(2.0)
    assert records[5][2]["linewidth"] == pytest.approx(0.7)
    assert records[6][2]["linewidth"] == pytest.approx(1.3)
    assert records[2][1] - records[1][1] == pytest.approx(0.26)
    assert records[0][2]["width"] == pytest.approx(0.46)
    assert records[1][2]["width"] == pytest.approx(0.30)
    assert records[2][2]["width"] == pytest.approx(0.16)
    assert records[1][2]["alpha"] == pytest.approx(0.42)
    assert records[1][2]["edge_alpha"] == pytest.approx(1.0)
    assert records[1][2]["median_linewidth"] == pytest.approx(1.5)
    assert records[2][2]["median_linewidth"] == pytest.approx(2.6)
    assert records[1][2]["median_width_fraction"] == pytest.approx(0.34)
    assert records[2][2]["median_width_fraction"] == pytest.approx(0.34)

    figure = captured_figures[0]
    assert figure.get_size_inches() == pytest.approx([12.4, 7.6])
    figure.canvas.draw()
    header = next(axis for axis in figure.axes if axis.get_gid() == "report-header")
    axis = next(axis for axis in figure.axes if axis.get_gid() != "report-header")
    assert header.texts[0].get_text() == (
        "Pressure distributions: overall, all-input, and favored subsets"
    )
    legend = header.get_legend()
    legend_labels = [text.get_text() for text in legend.get_texts()]
    assert legend_labels == [
        "All input",
        "Favored subset",
        "Selected model (>5%)",
    ]
    assert axis.spines["top"].get_visible()
    assert axis.spines["right"].get_visible()
    assert all(label.get_rotation() == 0 for label in axis.get_xticklabels())
    assert header.texts[0].get_fontweight() == "bold"
    assert legend.get_frame_on()
    assert axis.get_position().height > 0.65
    assert axis.get_position().y1 <= header.get_position().y0


def test_composition_plot_uses_fixed_limits_and_two_panels_without_liquid(
    monkeypatch: pytest.MonkeyPatch,
    independent_data: pd.DataFrame,
) -> None:
    captured_figures = []

    def capture_figure(fig, dpi=200):
        captured_figures.append(fig)
        return b"png"

    monkeypatch.setattr(excel_v2, "_figure_to_png", capture_figure)
    reference_df, kde_cache = excel_v2._resolve_reference_context(independent_data)
    excel_v2._build_model_selection_png(
        FakeWorkflow("P"),
        make_models("P"),
        make_composition_data(include_liquid=False),
        "P",
        reference_df,
        kde_cache,
        100,
    )

    figure = captured_figures[0]
    assert figure.get_size_inches() == pytest.approx([10.25, 5.1])
    figure.canvas.draw()
    header = next(axis for axis in figure.axes if axis.get_gid() == "report-header")
    plot_axes = [axis for axis in figure.axes if axis.get_gid() != "report-header"]
    assert header.texts[0].get_text() == (
        "Selected models across compositional space relative to independent experiments"
    )
    assert len(plot_axes) == 2
    assert header.get_legend() is not None
    for axis, spec in zip(plot_axes, excel_v2.COMPOSITION_PANEL_SPECS[:2]):
        np.testing.assert_allclose(axis.get_xlim(), spec["xlim"])
        np.testing.assert_allclose(axis.get_ylim(), spec["ylim"])
        assert axis.get_xlabel() == spec["xlabel"]
        assert axis.xaxis.label.get_fontsize() == pytest.approx(14)
        reference_scatter = axis.collections[0]
        assert reference_scatter.get_sizes()[0] == pytest.approx(14)
        assert reference_scatter.get_alpha() == pytest.approx(0.85)
        assert any(
            np.asarray(collection.get_sizes()).size
            and collection.get_sizes()[0] == pytest.approx(24)
            for collection in axis.collections
            if hasattr(collection, "get_sizes")
        )
    assert header.get_position().y0 - plot_axes[0].get_position().y1 > 0.01


def test_composition_plot_keeps_cpx_panel_pair_when_mg_number_is_missing(
    monkeypatch: pytest.MonkeyPatch,
    independent_data: pd.DataFrame,
) -> None:
    captured_figures = []

    def capture_figure(fig, dpi=200):
        captured_figures.append(fig)
        return b"png"

    source = make_composition_data(include_liquid=False).drop(
        columns=["MgO_cpx", "FeO_cpx"],
    )
    monkeypatch.setattr(excel_v2, "_figure_to_png", capture_figure)
    reference_df, kde_cache = excel_v2._resolve_reference_context(independent_data)
    excel_v2._build_model_selection_png(
        FakeWorkflow("P"),
        make_models("P"),
        source,
        "P",
        reference_df,
        kde_cache,
        100,
    )

    figure = captured_figures[0]
    figure.canvas.draw()
    plot_axes = [
        axis for axis in figure.axes if axis.get_gid() != "report-header"
    ]
    assert len(plot_axes) == 2
    assert any(
        text.get_text() == "Input values unavailable"
        for text in plot_axes[0].texts
    )
    assert not any(
        text.get_text() == "Input values unavailable"
        for text in plot_axes[1].texts
    )


def test_composition_plot_uses_precalculated_mg_number_alias() -> None:
    source = make_composition_data(include_liquid=False).drop(
        columns=["MgO_cpx", "FeO_cpx"],
    )
    source["Mg_number_cpx"] = 80.0

    _, mg_number = excel_v2._panel_xy(source, "cpx_cao_mg_number")

    assert mg_number is not None
    np.testing.assert_allclose(mg_number, 0.8)


def test_composition_plot_uses_four_panels_with_liquid(
    monkeypatch: pytest.MonkeyPatch,
    independent_data: pd.DataFrame,
) -> None:
    captured_figures = []

    def capture_figure(fig, dpi=200):
        captured_figures.append(fig)
        return b"png"

    monkeypatch.setattr(excel_v2, "_figure_to_png", capture_figure)
    reference_df, kde_cache = excel_v2._resolve_reference_context(independent_data)
    excel_v2._build_model_selection_png(
        FakeWorkflow("P"),
        make_models("P"),
        make_composition_data(include_liquid=True),
        "P",
        reference_df,
        kde_cache,
        100,
    )

    figure = captured_figures[0]
    plot_axes = [
        axis for axis in figure.axes if axis.get_gid() != "report-header"
    ]
    assert figure.get_size_inches() == pytest.approx([10.25, 9.1])
    assert len(plot_axes) == 4


def test_deviation_plot_has_english_title_and_pressure_tick_precision(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    captured_figures = []

    def capture_figure(fig, dpi=200):
        captured_figures.append(fig)
        return b"png"

    monkeypatch.setattr(excel_v2, "_figure_to_png", capture_figure)
    excel_v2._build_deviation_violin_png(
        FakeWorkflow("P"), make_models("P"), "P", 100
    )

    axis = captured_figures[0].axes[0]
    assert axis.get_title() == "Predicted deviation distributions by model"
    assert axis.yaxis.get_major_formatter().format_data_short(1.25) == "1.2"


def test_abbreviation_collision_adds_phase_suffix_only_for_plots() -> None:
    labels = model_abbreviation_lookup([MODEL_A, MODEL_B], "P", disambiguate=True)

    assert labels[MODEL_A] == "Pet20-cpx"
    assert labels[MODEL_B] == "Pet20-liq"
    assert model_abbreviation_lookup(["No selected model"], "P") == {
        "No selected model": "No selected model"
    }


def test_same_phase_abbreviation_collision_uses_stable_numbers() -> None:
    table_labels, plot_labels = excel_v2._report_model_abbreviation_lookups(
        [PU08_A, PU08_B],
        "T",
    )

    assert table_labels == {
        PU08_A: "Pu08_32d-1",
        PU08_B: "Pu08_32d-2",
    }
    assert plot_labels == table_labels


def test_duplicate_putirka_labels_are_shared_by_summary_and_votes(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workflow = FakeWorkflow("T")
    workflow.prediction_df = workflow.prediction_df.iloc[:, :2].copy()
    workflow.prediction_df.columns = [PU08_A, PU08_B]
    workflow.calculated_deviation_df = (
        workflow.calculated_deviation_df.iloc[:, :2].copy()
    )
    workflow.calculated_deviation_df.columns = [PU08_A, PU08_B]
    workflow.ood_mask_df = workflow.ood_mask_df.iloc[:, :2].copy()
    workflow.ood_mask_df.columns = [PU08_A, PU08_B]
    workflow._selected = pd.Series([PU08_A] * 12 + [PU08_B] * 8)
    models = [FakeModel(PU08_A, "T"), FakeModel(PU08_B, "T")]
    monkeypatch.setattr(
        excel_v2,
        "_build_model_selection_png",
        lambda *_args, **_kwargs: None,
    )
    monkeypatch.setattr(
        excel_v2,
        "_build_prediction_distribution_png",
        lambda *_args, **_kwargs: None,
    )
    monkeypatch.setattr(
        excel_v2,
        "_build_deviation_violin_png",
        lambda *_args, **_kwargs: None,
    )

    payload = excel_v2.build_report_payload(
        workflow,
        models,
        make_composition_data(),
        independent_data=make_composition_data(24),
    )

    assert payload.excel_model_summary_df is not None
    assert payload.excel_model_votes_df is not None
    assert payload.excel_model_summary_df["Model_abbreviation"].tolist() == [
        "Pu08_32d-1",
        "Pu08_32d-2",
    ]
    assert payload.excel_model_votes_df.loc[
        payload.excel_model_votes_df["Model_name"].isin([PU08_A, PU08_B]),
        "Model_abbreviation",
    ].tolist() == ["Pu08_32d-1", "Pu08_32d-2"]


def test_excel_formulas_formats_images_and_non_overlap(
    pressure_payload: ReportPayload,
) -> None:
    stream = io.BytesIO()
    excel_v2.write_report_excel(pressure_payload, stream)
    stream.seek(0)
    workbook = load_workbook(stream, data_only=False)
    worksheet = workbook["Model Summary"]
    summary = pressure_payload.excel_model_summary_df
    votes = pressure_payload.excel_model_votes_df
    assert summary is not None
    assert votes is not None
    assert worksheet.freeze_panes is None
    assert worksheet.column_dimensions["D"].width == pytest.approx(26.71, abs=0.1)
    assert worksheet.column_dimensions["F"].width == pytest.approx(23.71, abs=0.1)

    header_lookup = {
        worksheet.cell(1, column).value: column
        for column in range(1, worksheet.max_column + 1)
    }
    assert worksheet.cell(2, header_lookup["P_min (kbar)"]).number_format == "0.0"
    assert (
        worksheet.cell(2, header_lookup["Mean_calculated_deviation"]).number_format
        == "0.0"
    )
    vote_header_row = len(summary) + 4
    vote_first_row = vote_header_row + 1
    vote_total_row = vote_header_row + len(votes)
    assert worksheet.cell(vote_first_row, 4).value == (
        f"=IFERROR(C{vote_first_row}/$C${vote_total_row},0)"
    )
    assert worksheet.cell(vote_first_row, 4).number_format == "0.0%"
    assert worksheet.cell(vote_total_row, 3).value == (
        f"=SUM(C{vote_first_row}:C{vote_total_row - 1})"
    )
    assert worksheet.cell(vote_total_row, 4).value == (
        f"=SUM(D{vote_first_row}:D{vote_total_row - 1})"
    )

    assert len(worksheet._images) == 3
    anchors = sorted(
        (
            image.anchor._from.row,
            image.anchor._from.col,
            image.anchor.to.row,
            image.anchor.to.col,
        )
        for image in worksheet._images
    )
    prediction_anchor = next(anchor for anchor in anchors if anchor[1] == 4)
    lower_anchors = [anchor for anchor in anchors if anchor[1] == 0]
    composition_anchor, deviation_anchor = lower_anchors
    assert prediction_anchor[0] == len(summary) + 3
    assert prediction_anchor[2] - prediction_anchor[0] >= 24
    assert composition_anchor[0] == vote_total_row + 1
    assert composition_anchor[2] - composition_anchor[0] == 20
    assert composition_anchor[3] <= prediction_anchor[1]
    assert deviation_anchor[0] == composition_anchor[2] + 1
    assert deviation_anchor[0] > prediction_anchor[2]

    drawing_namespaces = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    }
    with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as archive:
        drawing = ET.fromstring(archive.read("xl/drawings/drawing1.xml"))
        extents = drawing.findall(".//a:xfrm/a:ext", drawing_namespaces)
        image_names = sorted(
            name for name in archive.namelist() if name.startswith("xl/media/")
        )
        assert len(extents) == len(image_names) == 3
        for extent, image_name in zip(extents, image_names):
            image_width, image_height = excel_v2._png_dimensions(
                archive.read(image_name)
            )
            displayed_ratio = float(extent.attrib["cx"]) / float(extent.attrib["cy"])
            assert displayed_ratio == pytest.approx(
                image_width / image_height,
                rel=1e-6,
            )
        locks = drawing.findall(".//a:picLocks", drawing_namespaces)
        assert locks and all(lock.attrib.get("noChangeAspect") == "1" for lock in locks)

    results_sheet = workbook["Results"]
    ranking_sheet = workbook["Ranking Details"]
    assert results_sheet.freeze_panes is None
    assert ranking_sheet.freeze_panes is None
    assert results_sheet.sheet_view.showGridLines is False
    assert ranking_sheet.sheet_view.showGridLines is False
    assert results_sheet.sheet_view.zoomScale == 90
    assert ranking_sheet.sheet_view.zoomScale == 90
    assert len(results_sheet.conditional_formatting) >= 3
    assert ranking_sheet["B1"].fill.fgColor.rgb == "FF1F4E78"
    assert ranking_sheet.auto_filter.ref is not None


def test_new_and_legacy_workbook_non_summary_sheets_are_identical(
    monkeypatch: pytest.MonkeyPatch,
    pressure_payload: ReportPayload,
) -> None:
    legacy_payload = ReportPayload(
        results_sheet=pressure_payload.results_sheet,
        model_summary_df=pressure_payload.model_summary_df,
        model_votes_df=pressure_payload.model_votes_df,
        ranking_details_df=pressure_payload.ranking_details_df,
        violin_plot_png=None,
        target="P",
    )
    new_stream = io.BytesIO()
    old_stream = io.BytesIO()
    excel_v2.write_report_excel(pressure_payload, new_stream)
    excel.write_report_excel_legacy(legacy_payload, old_stream)
    new_stream.seek(0)
    old_stream.seek(0)

    new_results = pd.read_excel(
        new_stream, sheet_name="Results", header=[0, 1], index_col=0
    )
    old_results = pd.read_excel(
        old_stream, sheet_name="Results", header=[0, 1], index_col=0
    )
    assert_frame_equal(new_results, old_results)
    new_stream.seek(0)
    old_stream.seek(0)
    new_ranking = pd.read_excel(
        new_stream, sheet_name="Ranking Details", index_col=0
    )
    old_ranking = pd.read_excel(
        old_stream, sheet_name="Ranking Details", index_col=0
    )
    assert_frame_equal(new_ranking, old_ranking)
