from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import Workbook, load_workbook

from aims4pt.reporting.excel import build_report_payload, write_report_excel
from web.services.input_validation import (
    CORE_CPX_COLUMNS,
    EXPECTED_INPUT_COLUMNS,
    LIQUID_OXIDE_COLUMNS,
    validate_workbook_bytes,
)


def _base_row() -> dict:
    values = {
        "Sample_ID": "sample-1",
        "SiO2_cpx": 51.2,
        "TiO2_cpx": 0.7,
        "Al2O3_cpx": 5.1,
        "FeOt_cpx": 6.8,
        "MnO_cpx": 0.2,
        "MgO_cpx": 15.1,
        "CaO_cpx": 21.0,
        "Na2O_cpx": 0.35,
        "K2O_cpx": 0.01,
        "Cr2O3_cpx": 0.25,
        "SiO2_liq": 50.1,
        "TiO2_liq": 1.1,
        "Al2O3_liq": 16.8,
        "FeOt_liq": 8.2,
        "MnO_liq": 0.15,
        "MgO_liq": 6.5,
        "CaO_liq": 9.8,
        "Na2O_liq": 3.2,
        "K2O_liq": 1.4,
        "Cr2O3_liq": 0.03,
        "P2O5_liq": 0.2,
        "H2O_liq": 1.0,
    }
    return values


def _workbook_bytes(
    *,
    columns: list[str] | None = None,
    rows: list[dict | None] | None = None,
    sheet_name: str = "compositions",
) -> bytes:
    columns = columns or EXPECTED_INPUT_COLUMNS
    rows = rows or [_base_row()]
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(["group"] * len(columns))
    worksheet.append(columns)
    for row in rows:
        if row is None:
            worksheet.append([None] * len(columns))
        else:
            worksheet.append([row.get(column) for column in columns])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _messages(result) -> list[str]:
    return [message.message for message in result.messages]


def test_valid_template_passes():
    result = validate_workbook_bytes(_workbook_bytes())

    assert result.is_valid
    assert not result.errors
    assert result.has_liquid
    assert result.cleaned_df is not None
    assert result.cleaned_df.loc[0, "Sample_ID"] == "sample-1"


def test_missing_compositions_sheet_gives_error():
    result = validate_workbook_bytes(_workbook_bytes(sheet_name="wrong"))

    assert not result.is_valid
    assert 'The sheet "compositions" is missing.' in _messages(result)


def test_missing_sample_id_gives_error():
    columns = [column for column in EXPECTED_INPUT_COLUMNS if column != "Sample_ID"]
    result = validate_workbook_bytes(_workbook_bytes(columns=columns))

    assert not result.is_valid
    assert "The Sample_ID column is missing." in _messages(result)


def test_missing_core_cpx_column_gives_error():
    missing = CORE_CPX_COLUMNS[0]
    columns = [column for column in EXPECTED_INPUT_COLUMNS if column != missing]
    result = validate_workbook_bytes(_workbook_bytes(columns=columns))

    assert not result.is_valid
    assert "One or more clinopyroxene composition columns are missing." in _messages(result)
    assert missing in result.errors[-1].columns


def test_non_float_composition_value_gives_error():
    row = _base_row()
    row["SiO2_cpx"] = "not-a-number"
    result = validate_workbook_bytes(_workbook_bytes(rows=[row]))

    assert not result.is_valid
    assert (
        "A composition column contains values that cannot be converted to float."
        in _messages(result)
    )


def test_blank_sample_id_gives_error():
    row = _base_row()
    row["Sample_ID"] = " "
    result = validate_workbook_bytes(_workbook_bytes(rows=[row]))

    assert not result.is_valid
    assert "Sample_ID contains blank values." in _messages(result)


def test_completely_empty_row_gives_error():
    result = validate_workbook_bytes(_workbook_bytes(rows=[_base_row(), None, _base_row()]))

    assert not result.is_valid
    assert "There are completely empty rows inside the data table." in _messages(result)


def test_duplicate_sample_id_gives_warning_only():
    row_a = _base_row()
    row_b = _base_row()
    result = validate_workbook_bytes(_workbook_bytes(rows=[row_a, row_b]))

    assert result.is_valid
    assert not result.errors
    assert "Duplicated Sample_ID values were found." in _messages(result)


def test_missing_oxide_values_are_filled_with_zero_and_warn():
    row = _base_row()
    row["TiO2_cpx"] = None
    result = validate_workbook_bytes(_workbook_bytes(rows=[row]))

    assert result.is_valid
    assert "Some missing oxide values were filled with 0." in _messages(result)
    assert result.cleaned_df.loc[0, "TiO2_cpx"] == 0


def test_missing_h2o_liq_is_filled_with_zero_and_warn():
    row = _base_row()
    row["H2O_liq"] = None
    result = validate_workbook_bytes(_workbook_bytes(rows=[row]))

    assert result.is_valid
    assert "Missing H2O_liq values were filled with 0." in _messages(result)
    assert result.cleaned_df.loc[0, "H2O_liq"] == 0


def test_no_liquid_data_disables_cpx_liq_but_allows_cpx_only():
    row = _base_row()
    for column in LIQUID_OXIDE_COLUMNS:
        row[column] = None
    result = validate_workbook_bytes(_workbook_bytes(rows=[row]))

    assert result.is_valid
    assert not result.errors
    assert not result.has_liquid
    assert (
        "Clinopyroxene-liquid calculations are unavailable because no liquid composition was provided."
        in _messages(result)
    )


def test_report_payload_can_write_non_empty_excel_report():
    import matplotlib

    matplotlib.use("Agg", force=True)

    class FakeModel:
        model_name = "Fake model"
        T_P = "P"
        X_cpx_training = pd.DataFrame({"P_kbar": [1.0, 2.0]})
        X_cpx_all = pd.DataFrame({"P_kbar": [1.0, 2.0]})
        rock_types = None
        uncertainty = 0.25
        y_min = 1.0
        y_max = 2.0

        def export_composition_range(self, *_args):
            return "SiO2_cpx: 50-52"

    class FakeWorkflow:
        T_P = "P"
        prediction_df = pd.DataFrame({"Fake model": [1.1, 1.8]})
        calculated_deviation_df = pd.DataFrame({"Fake model": [0.1, 0.2]})
        ood_mask_df = pd.DataFrame({"Fake model": [False, False]})
        failure_reason_df = pd.DataFrame({"Fake model": [0.1, 0.2]})

        def get_best_model_series(self):
            return pd.Series(["Fake model", "Fake model"])

    original = pd.DataFrame({"Sample_ID": ["a", "b"], "SiO2_cpx": [50, 51]})
    payload = build_report_payload(FakeWorkflow(), [FakeModel()], original)
    output = BytesIO()
    write_report_excel(payload, output)
    output.seek(0)

    workbook = load_workbook(output, read_only=True)
    assert set(workbook.sheetnames) == {"Results", "Model Summary", "Ranking Details"}
